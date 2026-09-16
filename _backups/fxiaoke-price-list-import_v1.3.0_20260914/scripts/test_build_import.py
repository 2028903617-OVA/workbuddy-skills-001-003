#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""build_import.py 的行为测试（TDD）。

运行： python scripts/test_build_import.py
原则：先让这些用例变绿，才允许改 build_import.py 的实现。

覆盖两个核心口径（2026-09-14 用户确认）：
  A. 组织(org) 决定「价目表」列；品牌 是产品实际品牌，与组织无关。
     - 欣暖家组织 → 美的价目表20250901；产品 TCL 彩电 → 品牌列 TCL
     - 乾鑫组织   → 乾鑫价目表20260701；产品 美的空调 → 品牌列 美的
     - 西门子组织 → 西门子价目表20260701；产品 西门子洗碗机 → 品牌列 西门子
  B. 品牌词表需覆盖 COLMO/东芝/索尼/松下/卡萨帝/科沃斯/长虹/容声/史密斯 等。
"""
import os
import sys
import json
import shutil
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_import as bi  # noqa: E402


class TestBrandInference(unittest.TestCase):
    """品牌 = 产品实际品牌，从产品名/型号推断。"""

    def test_common_brands(self):
        cases = [
            ("TCL平板电视98C10L-A", "TCL"),
            ("美的3匹变频柜机KFR-72LW/FQ1-1", "美的"),
            ("小天鹅滚筒洗干一体机TD10V628PLUS", "小天鹅"),
            ("西门子嵌入式洗碗机SJ43EB24KC", "西门子"),
            ("格力内机GMV-NHDR80PS/Fb无电辅带水泵", "格力"),
            ("方太油烟机CXW-358-03-X1A-W", "方太"),
            ("大金空调RJSYQ140AAV", "大金"),
        ]
        for model, expect in cases:
            with self.subTest(model=model):
                self.assertEqual(bi.guess_brand_from_model(model), expect)

    def test_newly_added_brands(self):
        """2026-09-14 实测缺失的品牌，必须能推断出来。"""
        cases = [
            ("COLMO多开门冰箱CRBUF706N-X1熔幔岩", "COLMO"),
            ("东芝主机XCY-MGP2241HTN-C", "东芝"),
            ("东芝内机XMD-GP0271MHN-C", "东芝"),
            ("东芝线控器RBC-ASXG31N-C", "东芝"),
            ("索尼电视机98XR70M2", "索尼"),
            ("松下多门冰箱NR-F611GA1-H", "松下"),
            ("卡萨帝壁挂炉LL1PBD26-CECONPXGU1", "卡萨帝"),
            ("科沃斯T90PRO水箱板白色", "科沃斯"),
            ("长虹电视机75Q70S", "长虹"),
            ("容声多门冰箱BCD-505Q50CZLBD", "容声"),
            ("史密斯净水机WF30D1", "史密斯"),
        ]
        for model, expect in cases:
            with self.subTest(model=model):
                self.assertEqual(bi.guess_brand_from_model(model), expect,
                                 f"品牌词表缺少 {expect}：{model}")

    def test_unknown_brand_returns_none(self):
        self.assertIsNone(bi.guess_brand_from_model("某某没听过的型号XYZ"))


class TestClassify(unittest.TestCase):
    """分类特例：乾鑫读源表、大金固定空调、欣暖家留空。"""

    def test_xinnuanjia_org_leaves_classification_blank(self):
        self.assertEqual(bi.classify("TCL平板电视98C10L-A", "欣暖家", "", ""),
                         ("", "", True))

    def test_dajin_fixed_aircon(self):
        self.assertEqual(bi.classify("大金某机型", "大金", "", ""),
                         ("空调", "空调", True))

    def test_qianxin_reads_source_columns(self):
        self.assertEqual(
            bi.classify("小天鹅1.5匹变频挂机", "乾鑫", "空调", "柜挂天井风管机家用多联"),
            ("空调", "柜挂天井风管机家用多联", True))

    def test_siemens_cross_fill(self):
        """西门子：大类列填【西门子品项】组，品项列填【西门子大类】组。"""
        pclass, pitem, matched = bi.classify("西门子微蒸烤一体机CP1K4R7T7W", "西门子", "", "")
        self.assertTrue(matched)
        self.assertEqual(pclass, "微蒸烤【西门子品项】")
        self.assertEqual(pitem, "微蒸烤箱【西门子大类】")


class TestRoundPrice(unittest.TestCase):
    def test_round_half_up(self):
        self.assertEqual(bi.round_price(423.52), 424)
        self.assertEqual(bi.round_price(2214.53), 2215)
        self.assertEqual(bi.round_price(14267.18), 14267)
        self.assertEqual(bi.round_price(6658.88888888889), 6659)
        self.assertEqual(bi.round_price(1.2), 1)
        self.assertEqual(bi.round_price(1.6), 2)

    def test_none_passthrough(self):
        self.assertIsNone(bi.round_price(None))


class TestGrade(unittest.TestCase):
    def test_single_letter(self):
        self.assertEqual(bi.extract_grade("D"), "D类")
        self.assertEqual(bi.extract_grade("a"), "A类")

    def test_empty_note(self):
        self.assertEqual(bi.extract_grade(None), "")
        self.assertEqual(bi.extract_grade(""), "")
        self.assertEqual(bi.extract_grade("2025-03-26 00:00:00"), "")


class TestProductMatching(unittest.TestCase):
    """产品匹配：不得把 F60-33Q7Pro(HE) 错配到 F60-33Q7Pro。"""

    PRODUCTS = [
        {"name": "美的电热水器F60-33Q7Pro", "mat": "美的电热水器F60-33Q7Pro",
         "spec": "F60-33Q7Pro", "code": "20.04.01.01.079",
         "pclass": "", "pitem": ""},
        {"name": "美的电热水器F60-33Q7Pro(HE)", "mat": "美的电热水器F60-33Q7Pro(HE)",
         "spec": "F60-33Q7Pro(HE)", "code": "20.04.01.01.099",
         "pclass": "", "pitem": ""},
    ]

    def test_he_suffix_not_mismatched(self):
        """源表要 (HE) 版本，绝不能返回不带 (HE) 的那条。"""
        name, how = bi.match_product("美的电热水器F60-33Q7Pro(HE)", self.PRODUCTS)
        self.assertEqual(name, "美的电热水器F60-33Q7Pro(HE)",
                         "错配到了不带 (HE) 的产品")

    def test_plain_version_still_matches(self):
        name, how = bi.match_product("美的电热水器F60-33Q7Pro", self.PRODUCTS)
        self.assertEqual(name, "美的电热水器F60-33Q7Pro")


class TestEndToEnd(unittest.TestCase):
    """端到端：组织决定价目表，品牌列是产品实际品牌。"""

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="plit_test_")
        self.src = os.path.join(self.tmp, "源表.xlsx")
        self.prod = os.path.join(self.tmp, "产品快照.xlsx")
        self.det = os.path.join(self.tmp, "明细快照.xlsx")
        self.outdir = os.path.join(self.tmp, "out")
        self._make_source()
        self._make_product_snapshot()
        self._make_detail_snapshot()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _make_source(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "申请表"
        ws.append(["纷享销客价目表新增（变更）申请表"])
        ws.append(["申请部门：欣暖家"])
        ws.append(["存货编码", "产品官方名称及型号", "零售建议价", "零售限制价",
                   "销售结算价", "实际成本价", "备注"])
        ws.append(["", "TCL平板电视98C10L-A", 21999, 14549, 12800, 12800, ""])
        ws.append(["", "COLMO多开门冰箱CRBUF706N-X1熔幔岩", 14400, 10666, 9600, 9600, ""])
        wb.save(self.src)

    def _make_product_snapshot(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "产品数据"
        ws.append(["产品名称（必填）", "物料名称", "型号规格", "物料编号",
                   "产品大类", "产品品相"])
        for n in ["TCL平板电视98C10L-A", "COLMO多开门冰箱CRBUF706N-X1熔幔岩"]:
            ws.append([n, n, n, "X", "", ""])
        wb.save(self.prod)

    def _make_detail_snapshot(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "价目表明细数据"
        ws.append(["唯一性ID（必填）", "价目表明细编号", "产品（必填）", "价目表（必填）",
                   "价目表_唯一性ID（必填）"])
        wb.save(self.det)

    def _run(self, org):
        import subprocess
        cmd = [sys.executable, os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                            "build_import.py"),
               "--source", self.src, "--org", org,
               "--product-snapshot", self.prod, "--detail-snapshot", self.det,
               "--outdir", self.outdir]
        r = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8")
        return r

    def _read_out(self):
        from openpyxl import load_workbook
        import glob
        fs = glob.glob(os.path.join(self.outdir, "*新建*.xlsx"))
        self.assertTrue(fs, "未生成新建文件")
        wb = load_workbook(fs[0], data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        hdr = [str(c) if c is not None else "" for c in rows[0]]
        return hdr, [dict(zip(hdr, r)) for r in rows[1:]]

    def test_xinnuanjia_org_uses_midea_pricebook_and_tcl_brand(self):
        r = self._run("欣暖家")
        self.assertEqual(r.returncode, 0, r.stderr)
        hdr, data = self._read_out()
        self.assertEqual(len(data), 2)
        by_prod = {d["产品（必填）"]: d for d in data}
        self.assertEqual(by_prod["TCL平板电视98C10L-A"]["品牌"], "TCL")
        self.assertEqual(by_prod["TCL平板电视98C10L-A"]["价目表（必填）"],
                         "美的价目表20250901")
        self.assertEqual(by_prod["COLMO多开门冰箱CRBUF706N-X1熔幔岩"]["品牌"], "COLMO")
        # 欣暖家组织 → 大类/品项留空
        self.assertIn(by_prod["TCL平板电视98C10L-A"]["产品大类"], (None, ""))
        self.assertIn(by_prod["TCL平板电视98C10L-A"]["产品品项"], (None, ""))

    def test_fixed_columns(self):
        self._run("欣暖家")
        hdr, data = self._read_out()
        for d in data:
            self.assertEqual(d["价目表折扣（%）"], 100)
            self.assertEqual(d["业务类型（必填）"], "预设业务类型")
            self.assertEqual(d["价目表售价"], 21999 if "TCL" in d["产品（必填）"] else 14400)

    def test_siemens_org_org_pricebook(self):
        """西门子组织 → 西门子价目表20260701。"""
        r = self._run("西门子")
        self.assertEqual(r.returncode, 0, r.stderr)
        hdr, data = self._read_out()
        for d in data:
            self.assertEqual(d["价目表（必填）"], "西门子价目表20260701")


if __name__ == "__main__":
    unittest.main(verbosity=2)
