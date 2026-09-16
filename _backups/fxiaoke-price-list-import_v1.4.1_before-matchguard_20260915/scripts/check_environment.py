#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""环境只读检查：Python 版本、openpyxl、两类快照可发现性。

输出 JSON: {"status": "ready|partial|needs_setup", "checks": [...]}
只读取状态，不安装、不联网、不修改任何配置。
"""
import glob
import json
import os
import sys

DOWNLOADS = r"D:\Backup\Downloads"


def find_snapshot(pattern, required_cols):
    files = sorted(glob.glob(os.path.join(DOWNLOADS, pattern)), reverse=True)
    for f in files:
        if os.path.basename(f).startswith("~$"):
            continue
        try:
            import openpyxl
            wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
            ws = wb[wb.sheetnames[0]]
            hdr = None
            for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
                hdr = [str(h) if h is not None else "" for h in row]
            ok = bool(hdr) and all(c in hdr for c in required_cols)
            if ok:
                for row in ws.iter_rows(min_row=2, max_row=2, values_only=True):
                    if not row or row[0] is None or str(row[0]).strip() == "":
                        ok = False
            wb.close()
            if ok:
                return f
        except Exception:
            continue
    return None


def main():
    checks = []

    ver = sys.version_info
    checks.append({"id": "python", "ok": ver >= (3, 10),
                   "detail": "Python %d.%d.%d" % (ver.major, ver.minor, ver.micro)})

    openpyxl_ok, openpyxl_detail = False, "未安装"
    try:
        import openpyxl
        openpyxl_ok = True
        openpyxl_detail = "openpyxl %s" % openpyxl.__version__
    except ImportError:
        pass
    checks.append({"id": "openpyxl", "ok": openpyxl_ok, "detail": openpyxl_detail})

    prod = find_snapshot("产品对象导出结果_*.xlsx", ["产品名称（必填）", "物料名称"])
    checks.append({"id": "product-snapshot", "ok": prod is not None,
                   "detail": prod or "Downloads 下找不到含必需列且有数据的产品对象导出"})

    det = find_snapshot("价目表明细对象导出结果_*.xlsx", ["产品（必填）", "价目表（必填）"])
    checks.append({"id": "detail-snapshot", "ok": det is not None,
                   "detail": det or "Downloads 下找不到含必需列且有数据的价目表明细导出"})

    if not openpyxl_ok:
        status = "needs_setup"
    elif not ver >= (3, 10):
        status = "needs_setup"
    elif prod is None and det is None:
        status = "partial"
    elif prod is None or det is None:
        status = "partial"
    else:
        status = "ready"

    print(json.dumps({"status": status, "checks": checks}, ensure_ascii=False, indent=2))
    sys.exit(0 if status in ("ready", "partial") else 1)


if __name__ == "__main__":
    main()
