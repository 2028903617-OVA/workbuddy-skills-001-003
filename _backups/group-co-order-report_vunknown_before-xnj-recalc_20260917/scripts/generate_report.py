# -*- coding: utf-8 -*-
"""
集团联单周报生成器 (group-co-order-report)

根据每周的"上周集团内部联单"Excel，生成符合规范的多工作表主文件、
按门店拆分的品牌子文件，以及按部门带单费汇总表。

用法:
    python generate_report.py \
        --src "D:/Backup/Downloads/上周集团内部联单_2026-07-15 14_12.xlsx" \
        --out "F:/集团联单数据/2026/7月/2026.07.06-07.12" \
        --week "2026.07.06-07.12" \
        [--ref "F:/集团联单数据/2026/4月/2026.04.13-04.19/集团联单明细2026.04.13-04.19.xlsx"] \
        [--remove-dajin]   # 临时去除大金公司数据

依赖: pandas, openpyxl
"""
import argparse
import os
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# === 默认参考文件 (含人员卡号370人 + 付款公司映射表) ===
DEFAULT_REF = r'F:/集团联单数据/2026/4月/2026.04.13-04.19/集团联单明细2026.04.13-04.19.xlsx'

# === 23列顺序 (前16列基础数据 → 姓名/银行卡号/开户行/付款公司 → 毛利+待回款) ===
COLS_23 = [
    '销售订单编号', '公司', '负责人主属部门', '负责人', '客户名称', '客户电话',
    '安装地址（老）', '销售订单金额(元)', '实际金额新', '内部员工主属品牌',
    '内部员工姓名', '现场管理/带单费', '内部员工带单提成系数', '带单类型', '客户来源',
    '确认时间', '姓名', '银行卡号', '开户行', '内部带单费付款公司',
    '销售结算毛利合计', '毛利合计', '待回款金额新'
]
# 19列 (去掉 姓名/银行卡号/开户行/付款公司, 第17-20位)
COLS_19 = COLS_23[:16] + COLS_23[20:]

# 金额列 (1-indexed)
MONEY_23 = [8, 9, 12, 13, 21, 22, 23]
SUM_23 = [8, 9, 12, 21, 22, 23]       # 内部员工带单提成系数(13) 不汇总
MONEY_19 = [8, 9, 12, 13, 17, 18, 19]
SUM_19 = [8, 9, 12, 17, 18, 19]

# 向后兼容别名
COLS_22 = COLS_23
COLS_18 = COLS_19
MONEY_22 = MONEY_23
SUM_22 = SUM_23
MONEY_18 = MONEY_19
SUM_18 = SUM_19

# L列 = 现场管理/带单费 = 第12列，小计/合计行仅此列标黄
L_COL = 12

# 品牌 Sheet 列表 (按公司列归类, 非内部员工主属品牌)
ALL_BRANDS = ['大金', '方太', '乾鑫', '西门子', '欣暖家', '启欣']

# 付款公司兜底规则: 参考表无该部门时，按公司列分配
PAY_FALLBACK = {
    '启欣': '杭州启欣智家电器有限公司',
    '乾鑫': '杭州乾鑫暖通设备工程有限公司',
    '方太': '杭州欣锋控股集团有限公司',
    '欣暖家': '杭州欣暖家节能科技有限公司',
    '西门子': '杭州欣锋控股集团有限公司',
}

# 不统计带单费的品牌: 内部员工主属品牌 属于这些值的整行直接删除
BRAND_EXCLUDE = ['泽锋']


def read_source(src, ref):
    df = pd.read_excel(src)
    df.columns = [str(c).strip().strip("'") for c in df.columns]

    ref_xls = pd.ExcelFile(ref)
    card_df = ref_xls.parse('集团公司人员卡号')
    card_map = {}
    for _, r in card_df.iterrows():
        name = str(r['姓名']).strip()
        card_map[name] = (str(r.get('银行卡号', '')), str(r.get('开户行', '')))

    pay_df = ref_xls.parse('内部带单费付款公司汇总表')
    pay_map = {}
    for _, r in pay_df.iterrows():
        dept = str(r['门店']).strip()
        pay_map[dept] = str(r.get('带单费付款公司选择', ''))

    return df, card_map, pay_map


import re

# 组号后缀模式：匹配末尾的 "N组" / "第N组"（如 1组、2组、第1组、第2组）
_GROUP_SUFFIX_RE = re.compile(r'[第]?\d+组\s*$')


def _resolve_pay_company(dept, company, pay_map):
    """解析付款公司：精确匹配 → 剥离组号后缀再匹配 → 公司兜底。"""
    # 1) 精确匹配
    if dept in pay_map and pay_map[dept]:
        return pay_map[dept]
    # 2) 剥离 "1组"/"2组"/"第1组" 等后缀，用基础部门名再查
    base = _GROUP_SUFFIX_RE.sub('', dept).strip()
    if base and base in pay_map and pay_map[base]:
        return pay_map[base]
    # 3) 公司级兜底
    return PAY_FALLBACK.get(company, '')


def fill_info(df, card_map, pay_map):
    """填充姓名/银行卡号/开户行/付款公司，打印匹配情况。"""
    df['姓名'] = ''
    df['银行卡号'] = ''
    df['开户行'] = ''
    df['内部带单费付款公司'] = ''

    unmatched_card = []
    unmatched_pay = []

    for i, r in df.iterrows():
        emp = str(r['内部员工姓名']).strip()
        dept = str(r['负责人主属部门']).strip()
        company = str(r['公司']).strip()
        if emp in card_map:
            df.at[i, '姓名'] = emp
            df.at[i, '银行卡号'] = card_map[emp][0]
            df.at[i, '开户行'] = card_map[emp][1]
        df.at[i, '内部带单费付款公司'] = _resolve_pay_company(dept, company, pay_map)

        status = 'OK' if (df.at[i, '银行卡号'] and df.at[i, '银行卡号'] != 'nan') else '未匹配'
        pay = df.at[i, '内部带单费付款公司']
        print(f"  {emp} | 卡号: {df.at[i, '银行卡号']} | 付款公司: {pay} | {status}")
        if status == '未匹配':
            unmatched_card.append(emp)
        if not pay or pay == 'nan':
            unmatched_pay.append((emp, dept))

    if unmatched_card:
        print(f"\n⚠️ 未匹配银行卡号: {unmatched_card}")
    if unmatched_pay:
        print(f"\n⚠️ 未匹配付款公司: {unmatched_pay}")
    return df, unmatched_card


# === 格式工具 ===
HEADER_FONT = Font(name='Arial', bold=True, color='FFFFFF', size=10)
HEADER_FILL = PatternFill('solid', fgColor='4472C4')
SUMMARY_FILL = PatternFill('solid', fgColor='FFFF00')
SUMMARY_FONT = Font(name='Arial', bold=True, size=10)
NORMAL_FONT = Font(name='Arial', size=10)
THIN_BORDER = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)
CENTER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)


def format_sheet(ws, ncols, money_indices, has_summary=True, sum_indices=None):
    if sum_indices is None:
        sum_indices = money_indices
    for col_idx in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER
        col_letter = get_column_letter(col_idx)
        max_w = len(str(cell.value or '')) * 2 + 4
        for row in range(2, ws.max_row + 1):
            c = ws.cell(row=row, column=col_idx)
            c.font = NORMAL_FONT
            c.border = THIN_BORDER
            c.alignment = Alignment(vertical='center', wrap_text=True)
            val_len = len(str(c.value or '')) * 1.2 + 4
            max_w = max(max_w, val_len)
        ws.column_dimensions[col_letter].width = min(max_w, 40)
        if col_idx in money_indices:
            for row in range(2, ws.max_row + 1):
                c = ws.cell(row=row, column=col_idx)
                if c.value is not None and c.value != '':
                    try:
                        float(c.value)
                        c.number_format = '#,##0.00'
                    except Exception:
                        pass

    if has_summary:
        sr = ws.max_row
        for col_idx in range(1, ncols + 1):
            c = ws.cell(row=sr, column=col_idx)
            c.font = SUMMARY_FONT
            if col_idx == L_COL:                       # 仅L列标黄
                c.fill = SUMMARY_FILL
            if col_idx in sum_indices:
                col_l = get_column_letter(col_idx)
                c.value = f'=SUM({col_l}2:{col_l}{sr-1})'
                c.number_format = '#,##0.00'
        ws.cell(row=sr, column=1).value = '合计'


def write_df_to_sheet(ws, df_data, cols, money_indices, has_summary=True, sum_indices=None):
    for ci, col_name in enumerate(cols, 1):
        ws.cell(row=1, column=ci, value=col_name)
    for ri, (_, row) in enumerate(df_data.iterrows()):
        for ci, col_name in enumerate(cols, 1):
            val = row.get(col_name, '')
            if pd.isna(val):
                val = ''
            ws.cell(row=ri + 2, column=ci, value=val)
    if has_summary and len(df_data) > 0:
        if sum_indices is None:
            sum_indices = money_indices
        sr = len(df_data) + 2
        for ci in range(1, len(cols) + 1):
            if ci in sum_indices:
                cl = get_column_letter(ci)
                ws.cell(row=sr, column=ci, value=f'=SUM({cl}2:{cl}{sr-1})')
        ws.cell(row=sr, column=1, value='合计')
    format_sheet(ws, len(cols), money_indices, has_summary and len(df_data) > 0, sum_indices)


def write_subfile_with_subtotals(ws, df_data, cols, money_indices, sum_indices, group_col='负责人主属部门'):
    """品牌子文件: 多部门时按部门加'小计'行 + 末尾'合计'行；单部门仅'合计'行。
    小计/合计行仅 L列(现场管理/带单费) 标黄，其余列只加粗不加底色。"""
    ncols = len(cols)
    for ci, col_name in enumerate(cols, 1):
        ws.cell(row=1, column=ci, value=col_name)

    current_row = 2
    subtotal_rows = []
    n_depts = df_data[group_col].nunique() if len(df_data) > 0 else 0

    if n_depts <= 1:
        for ri, (_, row) in enumerate(df_data.iterrows()):
            for ci, col_name in enumerate(cols, 1):
                val = row.get(col_name, '')
                if pd.isna(val):
                    val = ''
                ws.cell(row=ri + 2, column=ci, value=val)
        if len(df_data) > 0:
            sr = len(df_data) + 2
            for ci in sum_indices:
                cl = get_column_letter(ci)
                ws.cell(row=sr, column=ci, value=f'=SUM({cl}2:{cl}{sr-1})')
            ws.cell(row=sr, column=1, value='合计')
        format_sheet(ws, ncols, money_indices, len(df_data) > 0, sum_indices)
        return

    for dept, group_df in df_data.groupby(group_col, sort=False):
        group_start = current_row
        for _, row in group_df.iterrows():
            for ci, col_name in enumerate(cols, 1):
                val = row.get(col_name, '')
                if pd.isna(val):
                    val = ''
                ws.cell(row=current_row, column=ci, value=val)
            current_row += 1
        group_end = current_row - 1
        ws.cell(row=current_row, column=1, value=f'{dept} 小计')
        for ci in sum_indices:
            cl = get_column_letter(ci)
            ws.cell(row=current_row, column=ci,
                    value=f'=SUM({cl}{group_start}:{cl}{group_end})')
            ws.cell(row=current_row, column=ci).number_format = '#,##0.00'
        subtotal_rows.append(current_row)
        current_row += 1

    grand_total_row = current_row
    ws.cell(row=grand_total_row, column=1, value='合计')
    for ci in sum_indices:
        cl = get_column_letter(ci)
        formula = '=' + '+'.join([f'{cl}{r}' for r in subtotal_rows])
        ws.cell(row=grand_total_row, column=ci, value=formula)
        ws.cell(row=grand_total_row, column=ci).number_format = '#,##0.00'

    # 统一格式 (表头蓝底白字, 框线, 金额千分位)
    for col_idx in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER
        col_letter = get_column_letter(col_idx)
        max_w = len(str(cell.value or '')) * 2 + 4
        for row in range(2, grand_total_row + 1):
            c = ws.cell(row=row, column=col_idx)
            c.font = NORMAL_FONT
            c.border = THIN_BORDER
            c.alignment = Alignment(vertical='center', wrap_text=True)
            val_len = len(str(c.value or '')) * 1.2 + 4
            max_w = max(max_w, val_len)
        ws.column_dimensions[col_letter].width = min(max_w, 40)
        if col_idx in money_indices:
            for row in range(2, grand_total_row + 1):
                c = ws.cell(row=row, column=col_idx)
                if c.value is not None and c.value != '' and not str(c.value).startswith('='):
                    try:
                        float(c.value)
                        c.number_format = '#,##0.00'
                    except Exception:
                        pass

    # 仅 L列(现场管理/带单费) 标黄 + 加粗
    for sr in subtotal_rows + [grand_total_row]:
        for col_idx in range(1, ncols + 1):
            c = ws.cell(row=sr, column=col_idx)
            c.font = SUMMARY_FONT
            if col_idx == L_COL:
                c.fill = SUMMARY_FILL


def get_subfile_name(company, dept, pay_map):
    pmt = pay_map.get(dept, '')
    if '乾鑫' in pmt:
        suffix = '乾鑫'
    elif '智好' in pmt:
        suffix = '智好'
    elif '智德' in pmt:
        suffix = '智德'
    elif '欣锋' in pmt or '欣锋控股' in pmt:
        suffix = '欣锋'
    elif '启欣' in pmt:
        suffix = '启欣'
    else:
        suffix = company
    if company in suffix or suffix in company:
        label = company
    else:
        label = f'{company}（{suffix}）'
    return f'集团联单明细{WEEK_LABEL} - {label}.xlsx'


def build(src, out_dir, week_label, ref=DEFAULT_REF, remove_dajin=False):
    global WEEK_LABEL
    WEEK_LABEL = week_label
    os.makedirs(out_dir, exist_ok=True)

    df, card_map, pay_map = read_source(src, ref)
    print(f'原始数据: {len(df)} 条')
    print('公司分布:', df['公司'].value_counts().to_dict())
    print('负责人主属部门:', df['负责人主属部门'].value_counts().to_dict())
    print('\n=== 银行信息匹配情况 ===')
    df, unmatched = fill_info(df, card_map, pay_map)

    if remove_dajin:
        removed = df[df['公司'] == '大金']
        df = df[df['公司'] != '大金'].copy()
        print(f'\n已过滤大金数据 {len(removed)} 条，剩余 {len(df)} 条')

    # 剔除不统计带单费的品牌 (内部员工主属品牌 命中 BRAND_EXCLUDE)
    if BRAND_EXCLUDE:
        mask = df['内部员工主属品牌'].astype(str).str.strip().isin(BRAND_EXCLUDE)
        removed_brand = df[mask]
        if len(removed_brand) > 0:
            print(f'\n已剔除不统计带单费品牌({BRAND_EXCLUDE}) {len(removed_brand)} 条:')
            for _, r in removed_brand.iterrows():
                print(f"  - {r['内部员工姓名']} | {r['公司']} | {r['负责人主属部门']} | 品牌={r['内部员工主属品牌']}")
        df = df[~mask].copy()
        print(f'剔除后剩余 {len(df)} 条')

    # 排序: 公司 → 负责人主属部门
    df_sorted = df.sort_values(['公司', '负责人主属部门']).reset_index(drop=True)

    # === 主文件 ===
    wb = Workbook()
    wb.remove(wb.active)

    ws1 = wb.create_sheet('销售订单数据')
    write_df_to_sheet(ws1, df_sorted, COLS_18, MONEY_18, sum_indices=SUM_18)

    ws2 = wb.create_sheet('集团互联销售明细')
    write_df_to_sheet(ws2, df_sorted, COLS_22, MONEY_22, sum_indices=SUM_22)

    brands = [b for b in ALL_BRANDS if not (remove_dajin and b == '大金')]
    for brand in brands:
        ws_brand = wb.create_sheet(brand)
        brand_df = df_sorted[df_sorted['公司'] == brand].copy()
        write_df_to_sheet(ws_brand, brand_df, COLS_22, MONEY_22, sum_indices=SUM_22)

    # 人员卡号 sheet
    card_df = pd.read_excel(ref, sheet_name='集团公司人员卡号')
    ws9 = wb.create_sheet('集团公司人员卡号')
    for ci, h in enumerate(['序号', '姓名', '银行卡号', '开户行'], 1):
        ws9.cell(row=1, column=ci, value=h)
        ws9.cell(row=1, column=ci).font = HEADER_FONT
        ws9.cell(row=1, column=ci).fill = HEADER_FILL
        ws9.cell(row=1, column=ci).border = THIN_BORDER
        ws9.cell(row=1, column=ci).alignment = CENTER_ALIGN
    for ri, (_, r) in enumerate(card_df.iterrows()):
        ws9.cell(row=ri + 2, column=1, value=ri + 1)
        ws9.cell(row=ri + 2, column=2, value=r['姓名'])
        ws9.cell(row=ri + 2, column=3, value=r.get('银行卡号', ''))
        ws9.cell(row=ri + 2, column=4, value=r.get('开户行', ''))
        for c in range(1, 5):
            ws9.cell(row=ri + 2, column=c).font = NORMAL_FONT
            ws9.cell(row=ri + 2, column=c).border = THIN_BORDER
    for ci, w in zip(range(1, 5), [6, 12, 24, 10]):
        ws9.column_dimensions[get_column_letter(ci)].width = w

    # 付款公司汇总表 sheet
    pay_df = pd.read_excel(ref, sheet_name='内部带单费付款公司汇总表')
    ws10 = wb.create_sheet('内部带单费付款公司汇总表')
    for ci, h in enumerate(['门店', '带单费付款公司选择'], 1):
        ws10.cell(row=1, column=ci, value=h)
        ws10.cell(row=1, column=ci).font = HEADER_FONT
        ws10.cell(row=1, column=ci).fill = HEADER_FILL
        ws10.cell(row=1, column=ci).border = THIN_BORDER
        ws10.cell(row=1, column=ci).alignment = CENTER_ALIGN
    for ri, (_, r) in enumerate(pay_df.iterrows()):
        ws10.cell(row=ri + 2, column=1, value=r['门店'])
        ws10.cell(row=ri + 2, column=2, value=r.get('带单费付款公司选择', ''))
        for ci in range(1, 3):
            ws10.cell(row=ri + 2, column=ci).font = NORMAL_FONT
            ws10.cell(row=ri + 2, column=ci).border = THIN_BORDER
    ws10.column_dimensions['A'].width = 28
    ws10.column_dimensions['B'].width = 32

    main_path = f'{out_dir}/集团联单明细{week_label}.xlsx'
    wb.save(main_path)
    print(f'\nSaved main: {main_path}')

    # === 品牌子文件 (按 公司+部门 拆分, 含部门小计) ===
    sub_groups = {}
    for _, r in df_sorted.iterrows():
        key = (str(r['公司']).strip(), str(r['负责人主属部门']).strip())
        sub_groups.setdefault(key, []).append(r)

    sub_by_name = {}
    for (company, dept), records in sub_groups.items():
        sub_name = get_subfile_name(company, dept, pay_map)
        sub_by_name.setdefault(sub_name, []).extend(records)

    for sub_name, records in sub_by_name.items():
        sub_df = pd.DataFrame(records).reset_index(drop=True)
        sub_path = f'{out_dir}/{sub_name}'
        sub_wb = Workbook()
        sub_ws = sub_wb.active
        sub_ws.title = 'Sheet1'
        write_subfile_with_subtotals(sub_ws, sub_df, COLS_22, MONEY_22, sum_indices=SUM_22)
        sub_wb.save(sub_path)
        print(f'Sub-file: {sub_name} ({len(sub_df)} 条)')

    # === 按部门带单费汇总 ===
    print('\n=== 带单费按部门汇总 ===')
    dept_sum = df_sorted.groupby('负责人主属部门').agg(
        记录数=('销售订单编号', 'count'),
        带单费合计=('现场管理/带单费', 'sum'),
        实际金额合计=('实际金额新', 'sum'),
        毛利合计=('毛利合计', 'sum')
    ).reset_index()
    print(dept_sum.to_string())

    dept_wb = Workbook()
    dept_ws = dept_wb.active
    dept_ws.title = '带单费按部门汇总'
    dept_cols = ['负责人主属部门', '记录数', '实际金额合计', '带单费合计', '毛利合计']
    dept_money = [3, 4, 5]
    for ci, h in enumerate(dept_cols, 1):
        dept_ws.cell(row=1, column=ci, value=h)
        dept_ws.cell(row=1, column=ci).font = HEADER_FONT
        dept_ws.cell(row=1, column=ci).fill = HEADER_FILL
        dept_ws.cell(row=1, column=ci).border = THIN_BORDER
        dept_ws.cell(row=1, column=ci).alignment = CENTER_ALIGN
    dept_sorted = dept_sum.sort_values('带单费合计', ascending=False).reset_index(drop=True)
    for ri, (_, r) in enumerate(dept_sorted.iterrows()):
        for ci, c in enumerate(dept_cols, 1):
            dept_ws.cell(row=ri + 2, column=ci, value=r[c])
    tr = len(dept_sorted) + 2
    dept_ws.cell(row=tr, column=1, value='合计')
    for ci in dept_money:
        cl = get_column_letter(ci)
        dept_ws.cell(row=tr, column=ci, value=f'=SUM({cl}2:{cl}{tr-1})')
    format_sheet(dept_ws, len(dept_cols), dept_money, has_summary=True)
    for ci, w in zip(range(1, 6), [30, 8, 16, 14, 12]):
        dept_ws.column_dimensions[get_column_letter(ci)].width = w
    dept_path = f'{out_dir}/带单费汇总_按部门_{week_label}.xlsx'
    dept_wb.save(dept_path)
    print(f'\nSaved dept summary: {dept_path}')
    print(f'\n⚠️ 需手动补填银行卡号: {unmatched}')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='集团联单周报生成器')
    parser.add_argument('--src', required=True, help='源Excel路径')
    parser.add_argument('--out', required=True, help='输出目录')
    parser.add_argument('--week', required=True, help='周标签, 如 2026.07.06-07.12')
    parser.add_argument('--ref', default=DEFAULT_REF, help='参考文件(含卡号/付款公司映射)')
    parser.add_argument('--remove-dajin', action='store_true', help='临时去除大金公司数据')
    args = parser.parse_args()
    build(args.src, args.out, args.week, args.ref, args.remove_dajin)
