#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""特别分支：批量新增物料 —— 由「对方型号」生成 规格型号(F) 与 金蝶物料名(G)。

适用：价目表大调整 / 一次性新增大量型号 / 源表型号在销客里不存在需要新建。
只读写本地 xlsx，不联网、不写纷享销客。

用法:
  python build_new_material.py --source <源表.xlsx> --sheet <sheet名>
      [--col-brand 1] [--col-industry 3] [--col-model 5] [--col-f 6] [--col-g 7]
      [--product-snapshot <产品快照.xlsx>]
      [--extra-sample-sheet "已新增价目表明细"] [--extra-sample-col 5]
      [--outdir <输出目录>] [--dry-run]

输出:
  <outdir>/<源表名>_已补FG_<YYYYMMDD>.xlsx   源表副本（F/G 已填）
  <outdir>/待确认清单_<YYYYMMDD>.xlsx        低置信行
  stdout 摘要 JSON
"""
import argparse
import datetime
import json
import os
import re
import sys
import difflib
from collections import Counter, defaultdict

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
except ImportError:  # pragma: no cover
    print(json.dumps({"error": "openpyxl 未安装"}, ensure_ascii=False))
    sys.exit(2)

# ---------------------------------------------------------------- 常量
BRANDS = sorted(["卡萨帝", "COLMO", "西门子", "美的", "海信", "创维", "小天鹅", "东芝",
                 "海尔", "科沃斯", "容声", "方太", "格力", "添可", "老板", "TCL"],
                key=len, reverse=True)

IND_KEY = {
    "冰箱": ["冰箱", "冰吧", "酒柜", "冰柜"],
    "洗护": ["洗衣机", "洗干", "干衣机", "衣物护理", "洗烘"],
    "电视机": ["平板电视", "电视"],
    "厨房大电": ["洗碗机", "蒸烤", "烤箱", "蒸箱", "微波炉", "油烟机", "灶具", "消毒柜", "集成灶"],
    "居家清洁": ["扫地机器人", "洗地机", "吸尘", "扫地", "擦窗"],
    "空调": ["空调", "挂机", "柜机", "风管机", "多联"],
    "热水器": ["热水器", "采暖炉"],
    "水健康": ["净水", "管线机", "软水", "前置"],
}

# 品牌前缀：'/' 后面必须紧跟中文，避免误伤 KFR-72LW/U61-1、H1210-MBLNE/QT98TU1
BRAND_PREFIX = re.compile(r"^\s*[A-Za-z][A-Za-z\.\-]*/[一-龥]+\s*")
PROMO = [r"\d+\s*英寸", r"\d+\s*吋", r"官方标配", r"大屏护眼", r"世界杯",
         r"更护眼的家庭影院", r"&nbsp;", r"【新品】", r"2026款"]
TV_COLOR = ["流砂锖", "枪色", "钛金灰", "曜石黑", "钢琴黑", "太空灰", "星空灰"]
TV_BRAND_WORDS = ["海信RGB激光电视", "海信激光电视", "海信", "创维", "TCL", "电视"]
# 电视型号码后面的系列后缀词（无数字，正则抽码时会漏掉，需补回）
SERIES_WORD = re.compile(r"^(Ultra|Pro|Max|Plus|Mini|SE|X)$", re.I)
# 型号码：允许数字开头（75E8S / 100L7QR）
TAIL = re.compile(r"[A-Za-z0-9][A-Za-z0-9][A-Za-z0-9\-\.\(\)\+/]{3,}\s*$")
# 品项尾部脏值：CE / CZ / CG / CEC 等系列码（2-4 个大写字母），如「滚筒洗衣机CEC」
JUNK_ITEM = re.compile(r"[A-Z]{2,4}\s*$")
BAD_ITEM = re.compile(r"IOT|智能物联|赠品|样机|空机$|套装$")
AC_P = {"25": "1", "26": "1", "32": "1.5", "33": "1.5", "35": "1.5", "36": "1.5",
        "40": "2", "45": "2", "46": "2", "50": "2", "51": "2", "60": "3", "61": "3",
        "71": "3", "72": "3", "73": "3", "88": "4", "120": "5"}


# ---------------------------------------------------------------- 基础工具
def industry_of(item):
    for ind, kws in IND_KEY.items():
        if any(k in item for k in kws):
            return ind
    return "其他"


def lcp(a, b):
    n = 0
    for x, y in zip(a.lower(), b.lower()):
        if x != y:
            break
        n += 1
    return n


def parse_name(name, brands=None):
    """'容声对开门冰箱BCD-606WKK1FPGZA' -> (容声, 对开门冰箱, BCD-606WKK1FPGZA)"""
    n = str(name or "").strip()
    if not n:
        return None
    br = None
    for b in (brands or BRANDS):
        if n.startswith(b):
            br, rest = b, n[len(b):]
            break
    if not br:
        return None
    m = TAIL.search(rest)
    if not m:
        return None
    item, code = rest[:m.start()].strip(), m.group(0).strip()
    item = re.sub(r"\d+\s*$", "", item).strip()   # 平板电视75 -> 平板电视
    item = JUNK_ITEM.sub("", item).strip()        # 滚筒洗衣机CE -> 滚筒洗衣机
    if not item:
        return None
    return (br, item, code)


def load_products(path):
    """产品快照 -> list of {name, mat, spec}"""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    hdr = [str(h).strip() if h is not None else ""
           for h in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]

    def ix(name):
        return hdr.index(name) if name in hdr else None
    i_n, i_m, i_s = ix("产品名称（必填）"), ix("物料名称"), ix("型号规格")
    out = []
    for r in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        def g(i):
            return (str(r[i]).strip() if i is not None and i < len(r) and r[i] is not None else "")
        if not g(i_n):
            continue
        out.append({"name": g(i_n), "mat": g(i_m), "spec": g(i_s)})
    wb.close()
    return out


def build_lib(products, extra_names=None):
    """-> {品牌: [(型号码, 品项), ...]}"""
    lib = defaultdict(list)
    for p in products:
        s = parse_name(p.get("name"))
        if s:
            lib[s[0]].append((s[2], s[1]))
    for n in (extra_names or []):
        s = parse_name(n)
        if s:
            lib[s[0]].append((s[2], s[1]))
    return lib


# ---------------------------------------------------------------- 清洗
def clean_model(e, industry=None):
    """E 列 -> 规格型号 F（通用，非电视）"""
    s = str(e or "").replace("&nbsp;", " ").replace("\xa0", " ").strip()
    s = re.sub(r"^\s*【新品】\s*", "", s)
    s = BRAND_PREFIX.sub("", s)
    s = re.sub(r"^\s*【新品】\s*", "", s)
    for pat in PROMO:
        s = re.sub(pat, "", s)
    s = re.sub(r"\s{2,}", " ", s).strip()
    # 只在「字母+连字符」后合空格：MR- 449 -> MR-449；不能误伤 芙万 Artist 50S
    s = re.sub(r"([A-Za-z]-)\s+(\d)", r"\1\2", s)
    half = len(s) // 2
    if len(s) > 8 and half > 3 and s[:half].strip() == s[len(s) - half:].strip():
        s = s[:half].strip()                                # 重复片段去重
    return s


def tv_model(e):
    """电视机专用：抽型号码，抽不到就退回清洗后的中文系列名"""
    s = re.sub(r"^\s*【新品】\s*", "", str(e or "")).replace("&nbsp;", " ").replace("\xa0", " ")
    s = BRAND_PREFIX.sub("", s)
    for pat in PROMO:
        s = re.sub(pat, "", s)
    for b in TV_BRAND_WORDS:
        s = s.replace(b, " ")
    for c in TV_COLOR:
        s = s.replace(c, " ")
    s = re.sub(r"\s{2,}", " ", s).strip()
    codes = [x for x in re.findall(r"[A-Za-z0-9][A-Za-z0-9\-]{4,}", s) if re.search(r"\d", x)]
    if codes:
        best = sorted(codes, key=len, reverse=True)[0]
        # 紧随型号码的系列后缀要保留：115Q10M Ultra / 98Q10M Pro（不能只留 115Q10M）
        after = s.split(best, 1)[1].split() if best in s else []
        if after and SERIES_WORD.match(after[0]):
            return "%s %s" % (best, after[0])
        return best
    return s


def primary_code(spec):
    toks = [x for x in re.findall(r"[A-Za-z0-9][A-Za-z0-9\-\.\(\)\+/]{3,}", spec)
            if re.search(r"\d", x)]
    if not toks:
        return spec[:14]
    return sorted(toks, key=len, reverse=True)[0].split("+")[0]


# ---------------------------------------------------------------- 品项推断
def ac_item(code):
    """KFR-35GW -> 1.5匹变频挂机"""
    m = re.search(r"KFR-(\d+)(GW|LW)", code, re.I)
    if not m:
        return "", 0.0, ""
    num, form = m.group(1), m.group(2).upper()
    pin = AC_P.get(num, "")
    typ = "变频挂机" if form == "GW" else "变频柜机"
    if not pin:
        return "", 0.0, ""
    return "%s匹%s" % (pin, typ), 0.7, "空调型号推导"


def guess_item(brand, code, industry, lib):
    """-> (品项, 置信度, 来源)"""
    cands = lib.get(brand, [])
    best = (0, 0.0, "")
    for c, item in cands:
        l = lcp(code, c)
        rt = difflib.SequenceMatcher(None, code.lower(), c.lower()).ratio()
        if (l, rt) > (best[0], best[1]):
            best = (l, rt, item)
    l, rt, item = best
    if item and industry and industry_of(item) != industry:
        # 行业不符：借来的品项直接丢弃（不能再走 型号相似 高置信返回），改走兜底
        item, l, rt = "", 0, 0.0
    if item and l >= 5:
        return item, round(min(0.99, 0.45 + l / 22.0 + rt / 5), 2), "最相近型号"
    if item and rt >= 0.75:
        return item, round(rt, 2), "型号相似"
    if industry == "空调":
        it, cf, src = ac_item(code)
        if it:
            return it, cf, src
    # 同品牌同行业众数 优先于 行业兜底（众数更贴近该品牌实际叫法）
    ctr = Counter(it for _, it in cands if industry_of(it) == industry)
    if ctr:
        it, v = ctr.most_common(1)[0]
        return it, round(v / sum(ctr.values()), 2), "行业众数 %d/%d" % (v, sum(ctr.values()))
    if industry == "水健康":
        return "净水机", 0.3, "水健康兜底"
    if industry == "热水器":
        return "燃气热水器", 0.3, "热水器兜底"
    return "", 0.0, "无可用样本"


def fix_item(item, industry, code, brand):
    """脏值过滤 + 西门子厨电型号特征"""
    if item and BAD_ITEM.search(item):
        item = ""
    if not item:
        if industry == "热水器":
            item = "燃气热水器"
        elif industry == "水健康":
            item = "净水机"
    if industry == "厨房大电" and brand == "西门子":
        if code.upper().startswith("CXW"):
            item = "国产欧式油烟机"
        elif "蒸烤" in code or re.match(r"^CS\d", code, re.I):
            item = "国产嵌入式蒸烤一体机"
    return item


def compose_g(brand, item, spec):
    """G = 品牌 + 品项 + 规格型号；spec 已含「品牌+品项」或已含品牌时不重复"""
    if not (item and spec):
        return ""
    if brand and brand in spec and item in spec:
        return spec                      # spec 已含 品牌+品项（如 海信激光电视星光S1）
    head = "" if spec.startswith(brand) else brand
    return "%s%s%s" % (head, item, spec)


# ---------------------------------------------------------------- 主流程
def process_rows(rows, lib):
    """rows: [{row, brand, industry, E}] -> 追加 F/G/item/conf/src"""
    out = []
    for r in rows:
        industry = r["industry"]
        if industry == "电视机":
            spec = tv_model(r["E"])
            item = "激光电视" if "激光电视" in str(r["E"]) else "平板电视"
            conf, src = 0.75, "电视规则"
        else:
            spec = clean_model(r["E"], industry)
            item, conf, src = guess_item(r["brand"], primary_code(spec), industry, lib)
            item = fix_item(item, industry, primary_code(spec), r["brand"])
        if not spec:
            conf = 0.0
        out.append(dict(r, F=spec, G=compose_g(r["brand"], item, spec),
                        item=item, conf=conf, src=src))
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--source", required=True)
    ap.add_argument("--sheet", required=True)
    ap.add_argument("--col-brand", type=int, default=1)
    ap.add_argument("--col-industry", type=int, default=3)
    ap.add_argument("--col-model", type=int, default=5, help="E 列（对方型号标题）")
    ap.add_argument("--col-f", type=int, default=6)
    ap.add_argument("--col-g", type=int, default=7)
    ap.add_argument("--header-row", type=int, default=1)
    ap.add_argument("--product-snapshot", default=None)
    ap.add_argument("--downloads", default=r"D:\Backup\Downloads")
    ap.add_argument("--extra-sample-sheet", default=None)
    ap.add_argument("--extra-sample-col", type=int, default=5)
    ap.add_argument("--outdir", default=None)
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    if not os.path.exists(args.source):
        print(json.dumps({"error": "源表不存在: %s" % args.source}, ensure_ascii=False))
        sys.exit(2)

    wb = openpyxl.load_workbook(args.source, data_only=True)
    if args.sheet not in wb.sheetnames:
        print(json.dumps({"error": "找不到 sheet: %s（现有: %s）"
                          % (args.sheet, wb.sheetnames)}, ensure_ascii=False))
        sys.exit(2)
    ws = wb[args.sheet]

    # 样本库
    snap = args.product_snapshot
    if not snap:
        import glob as _g
        cands = sorted(_g.glob(os.path.join(args.downloads, "产品对象导出结果_*.xlsx")), reverse=True)
        snap = next((f for f in cands if not os.path.basename(f).startswith("~$")), None)
    products = load_products(snap) if snap else []
    extra = []
    if args.extra_sample_sheet and args.extra_sample_sheet in wb.sheetnames:
        ws2 = wb[args.extra_sample_sheet]
        for r in ws2.iter_rows(min_row=args.header_row + 1, values_only=True):
            v = r[args.extra_sample_col - 1] if len(r) >= args.extra_sample_col else None
            if v:
                extra.append(str(v))
    lib = build_lib(products, extra)

    rows = []
    for ridx in range(args.header_row + 1, ws.max_row + 1):
        e = ws.cell(row=ridx, column=args.col_model).value
        d = ws.cell(row=ridx, column=args.col_model - 1).value
        if not str(e or "").strip() and not str(d or "").strip():
            continue
        rows.append({"row": ridx,
                     "brand": str(ws.cell(row=ridx, column=args.col_brand).value or "").split("/")[-1].strip(),
                     "industry": str(ws.cell(row=ridx, column=args.col_industry).value or "").strip(),
                     "E": str(e or "")})
    res = process_rows(rows, lib)

    today = datetime.date.today().strftime("%Y%m%d")
    outdir = args.outdir or os.path.dirname(os.path.abspath(args.source))
    os.makedirs(outdir, exist_ok=True)
    n_empty = sum(1 for r in res if not r["F"] or not r["G"])
    n_low = sum(1 for r in res if r["conf"] < 0.7)

    if not args.dry_run:
        for r in res:
            ws.cell(row=r["row"], column=args.col_f).value = r["F"]
            ws.cell(row=r["row"], column=args.col_g).value = r["G"]
        base = os.path.splitext(os.path.basename(args.source))[0]
        p1 = os.path.join(outdir, "%s_已补FG_%s.xlsx" % (base, today))
        wb.save(p1)

        wb2 = openpyxl.Workbook()
        ws2 = wb2.active
        ws2.title = "待确认清单"
        ws2.append(["低置信(置信<0.7)共 %d 行，请人工复核品项" % n_low])
        ws2.append(["行号", "品牌", "行业", "E列型号标题", "F规格型号", "G物料名称", "品项", "来源", "置信"])
        for r in res:
            if r["conf"] < 0.7:
                ws2.append([r["row"], r["brand"], r["industry"], r["E"], r["F"], r["G"],
                            r["item"], r["src"], r["conf"]])
        for i in range(1, 10):
            c = ws2.cell(row=2, column=i)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="2F5597")
        for i, w in enumerate([8, 12, 12, 40, 34, 44, 20, 18, 10], start=1):
            ws2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        p2 = os.path.join(outdir, "待确认清单_%s.xlsx" % today)
        wb2.save(p2)
    else:
        p1 = p2 = "(dry-run 未写文件)"

    print(json.dumps({
        "sheet": args.sheet, "数据行": len(res),
        "样本库": sum(len(v) for v in lib.values()), "品牌数": len(lib),
        "F或G为空": n_empty, "低置信": n_low,
        "品项来源": dict(Counter(r["src"] for r in res).most_common()),
        "files": [p1, p2],
    }, ensure_ascii=False, indent=2))
    wb.close()


if __name__ == "__main__":
    main()
