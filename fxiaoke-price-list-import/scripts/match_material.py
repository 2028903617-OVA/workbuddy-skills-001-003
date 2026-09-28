#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""能力 2：我方物料匹配外部物料。

外部拿来一份物料/型号清单，回答「这些在我们系统里对应哪个物料、有没有、销客产品编码是多少」。
只读写本地 xlsx，不联网、不调子进程、不写纷享销客。

用法:
  python match_material.py --source <外部物料表.xlsx> [--sheet <名>] [--org 天猫]
      [--material-table <金蝶物料表.xlsx> ...]        # 可多张，按「使用组织」列自动分流
      [--product-snapshot <产品快照.xlsx>]
      [--outdir <目录>] [--dry-run]

匹配顺序（按序短路，见 references/ability-2-material-match.md）：
  1 物料名称精确  →  2 规格型号精确  →  3 型号标题模糊（每个候选必须过两道闸）
铁律：能用「物料名称」列就绝不用「型号标题」列——型号标题与我方物料名零个完全一致。

实现注意：外部表与物料表都用 openpyxl **read_only 模式**读，read_only 下 `ws.cell()`
会反复重解析整表导致极慢，因此本脚本一律走 `iter_rows`，不调用 `ws.cell()`。

输出:
  <outdir>/外部物料匹配结果_{org}_{YYYYMMDD}.xlsx
    sheet 匹配结果 / 未匹配清单 / 统计
"""
import argparse
import datetime
import glob
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_import as bi  # noqa: E402  共用：表头语义识别、模糊匹配两道闸、快照加载

# 外部表列语义（命中其一即认；不按列字母定位）
EXT_SEMANTICS = {
    "model":    ["型号标题", "商品名称", "货品名称", "产品官方名称及型号", "型号"],
    "material": ["物料名称", "产品名称", "金蝶名称", "我方名称", "内部物料名称"],
    "spec":     ["规格型号", "规格"],
    "brand":    ["品牌", "品牌系"],
    "industry": ["行业", "品类", "产品大类", "大类"],
    "code":     ["商品编码", "商家编码", "SKU编码", "编码", "SKU", "店铺ID"],
}
SHORT_KW = 2
PLACEHOLDER = {"", "0", "0.0", "-", "无", "NA", "N/A", "NAN", "NONE"}
CONF = {"物料名称精确": 1.0, "规格型号精确": 0.95, "型号标题模糊": 0.8, "未匹配": 0.0}
ADD_COLS = ["我方产品名称", "我方物料名称", "物料编码", "规格型号", "销客产品编码",
            "所属组织", "匹配方式", "置信度", "来源", "备注"]


# ---------------------------------------------------------------- read_only 安全取数
def read_rows(ws, min_row=1, max_row=None):
    """返回 [(行号, [值...])]。统一走 iter_rows，避免 read_only 下 ws.cell() 的性能陷阱。"""
    out = []
    for i, row in enumerate(ws.iter_rows(min_row=min_row, max_row=max_row, values_only=True),
                            start=min_row):
        out.append((i, list(row)))
    return out


def row_at(rows, r):
    for i, vals in rows:
        if i == r:
            return vals
    return []


# 可当匹配主键的三类列：型号类 / 物料名称类 / 规格型号类。
# 不要求必须有型号类——金蝶供货价表这种源表只有「物料名称+规格型号」没有型号标题，
# 而按铁律它恰恰应该走第一级「物料名称精确」。
KEY_COLS = ("model", "material", "spec")


def find_header_row(rows, scan=8):
    """表头行 = 前 scan 行里命中语义最多、且含三类可匹配列之一的那一行。"""
    best, best_n = 1, -1
    for i, vals in rows[:scan]:
        hits = set()
        for key, kws in EXT_SEMANTICS.items():
            if any(bi.header_hit(c, kw) for c in vals for kw in kws):
                hits.add(key)
        if len(hits) > best_n and (set(KEY_COLS) & hits):
            best, best_n = i, len(hits)
    return best


def find_col(hdr, kws):
    for kw in kws:                                    # 一轮：整格相等
        for i, c in enumerate(hdr, start=1):
            if bi.norm_header(c) == kw:
                return i
    for kw in kws:                                    # 二轮：包含（仅长词）
        if len(kw) <= SHORT_KW:
            continue
        for i, c in enumerate(hdr, start=1):
            nc = bi.norm_header(c)
            if nc and kw in nc:
                return i
    return None


def resolve_columns(hdr):
    return {k: find_col(hdr, kws) for k, kws in EXT_SEMANTICS.items()}


def _last_content_col(rows, hdr_row):
    """最后一个「表头或数据里有内容」的列号（1 起）。用于裁掉空尾巴列。"""
    last = 0
    for i, vals in rows:
        if i < hdr_row:
            continue
        for j, v in enumerate(vals, start=1):
            if v is not None and str(v).strip() != "":
                last = max(last, j)
    return last


def pick_sheet(wb, hints=("货品", "物料", "型号", "商品", "价目", "匹配")):
    """选外部表的数据 sheet。
    候选门槛：三类可匹配列（型号/物料名称/规格型号）任有其一直接算；
    SheetN 名不硬跳——真数据也可能装在唯一一个 Sheet1 里（实测踩过），
    但多张带语义的候选并存时优先用非 SheetN 名的，SheetN 只在没得选时才用。"""
    named, sheetn = [], []
    for sn in wb.sheetnames:
        if sn.lower().startswith("hidden"):
            continue
        rows = read_rows(wb[sn], 1, 8)
        if not rows:
            continue
        hdr = row_at(rows, find_header_row(rows))
        if any(bi.header_hit(c, kw)
               for c in hdr for k in KEY_COLS for kw in EXT_SEMANTICS[k]):
            (sheetn if re.fullmatch(r"[Ss]heet\d+", sn) else named).append(sn)
    cands = named or sheetn
    if not cands:
        return None
    hinted = [s for s in cands if any(h in s for h in hints)]
    return (hinted or cands)[-1]


# ---------------------------------------------------------------- 我方底表
def load_material_table(path):
    """金蝶物料表 -> list of dict。列名按语义匹配，不按列字母。"""
    wb = bi.open_wb(path)
    ws = wb[wb.sheetnames[0]]
    rows = read_rows(ws, 1)
    if not rows:
        wb.close()
        return []
    hdr = rows[0][1]
    want = {"org": ["使用组织"], "code": ["编码", "物料编码"], "name": ["名称", "物料名称"],
            "spec": ["规格型号", "规格"], "xk": ["销客产品编码"], "group": ["物料分组"]}
    ix = {k: find_col(hdr, v) for k, v in want.items()}
    out = []
    for _, r in rows[1:]:

        def g(key):
            i = ix[key]                                # find_col 返回 1 起的列号
            return (str(r[i - 1]).strip() if i and i <= len(r) and r[i - 1] is not None else "")

        if not g("name"):
            continue
        out.append({"src": os.path.basename(path), "org": g("org"), "code": g("code"),
                    "name": g("name"), "spec": g("spec"), "xk": g("xk"), "group": g("group")})
    wb.close()
    return out


def default_material_tables(downloads):
    fs = sorted(glob.glob(os.path.join(downloads, "物料_*.xlsx")), reverse=True)
    return [f for f in fs if not os.path.basename(f).startswith("~$")]


# ---------------------------------------------------------------- 索引与匹配
def build_index(materials, products, org=None):
    """建三级匹配用的索引，避免逐行扫全表。

    ⚠️ 两个来源必须都进索引，它们的「名称」不是一回事（实测踩过）：
      产品主数据快照「物料名称」→ 外部表的物料名称列与它几乎 100% 完全一致；
      金蝶物料表「名称」      → 覆盖范围不同（更旧的子集），但**只有它带「销客产品编码」**。
    所以：先用快照命中行，再回物料表按名称/规格补销客产品编码。
    """
    pool = [m for m in materials if not org or m["org"] == org] or materials
    by_name, by_spec, spec_pairs = {}, {}, []
    for m in pool:
        if m["name"]:
            by_name.setdefault(m["name"], []).append(m)
        if m["spec"]:
            by_spec.setdefault(m["spec"], []).append(m)
            if len(m["spec"]) >= 4:
                spec_pairs.append((m["spec"], m))
    prod_by_mat, prod_by_name, prod_by_spec = {}, {}, {}
    for p in products:
        if p.get("mat"):
            prod_by_mat.setdefault(p["mat"], p)
        if p.get("name"):
            prod_by_name.setdefault(p["name"], p)
        if p.get("spec"):
            prod_by_spec.setdefault(p["spec"], p)
    return {"pool": pool, "by_name": by_name, "by_spec": by_spec, "spec_pairs": spec_pairs,
            "prod_by_mat": prod_by_mat, "prod_by_name": prod_by_name, "prod_by_spec": prod_by_spec,
            "has_products": bool(products)}


class Matcher:
    """按组织缓存索引；组织不传则跨组织匹配。"""

    def __init__(self, materials, products=None):
        self.materials = materials
        self.products = products or []
        self._cache = {}

    def index(self, org=None):
        key = org or ""
        if key not in self._cache:
            self._cache[key] = build_index(self.materials, self.products, org)
        return self._cache[key]


def _xk_of(hits):
    """同名/同规格多条：销客编码唯一才算同一条物料。"""
    xs = {h["xk"] for h in hits if h["xk"]}
    return (list(xs)[0] if len(xs) == 1 else "") if xs else ""


def _pack_mat(m, method, note=""):
    # 物料表这条本身没销客产品编码（且快照也没覆盖该物料，否则会先走 _pack_prod）→ 写明，别静默留空
    if not m["xk"] and not note:
        note = "金蝶物料表该物料无销客产品编码，快照也未覆盖，待人工确认"
    return {"我方产品名称": "", "我方物料名称": m["name"], "物料编码": m["code"],
            "规格型号": m["spec"], "销客产品编码": m["xk"], "所属组织": m["org"],
            "匹配方式": method, "置信度": CONF[method], "来源": m["src"], "备注": note}


def _pack_prod(p, method, idx):
    """命中产品主数据快照时的打包。

    「销客产品编码」的取值顺序（不猜，取不到就留空并写备注）：
      1 金蝶物料表按「物料名称 → 规格型号」查到唯一编码——官方维护的对照关系，最权威；
      2 物料表没覆盖该物料 → 回落用快照自己的「唯一性ID（必填）」。
        依据：2026-09-22 交叉验证，快照唯一性ID 与物料表销客产品编码逐行一致 1103/1104，
        且它本来就是该产品在纷享销客里的记录 ID。
    快照「产品编码」列实测大面积为空、「物料编号」是 ERP 存货编码，两者都不能当销客编码用。
    """
    xk, note = "", ""
    for key, table in ((p["mat"], "by_name"), (p["spec"], "by_spec")):
        if not key:
            continue
        one = _xk_of(idx[table].get(key, []))
        if one:
            xk = one
            break
        # 该名称在物料表里有多条且编码不一致 → 不猜，留空并记备注
        hits = idx[table].get(key, [])
        if len(hits) > 1 and not _xk_of(hits):
            note = "物料表同名多条、编码不一致，销客产品编码待人工确认"
    if not xk and not note:
        uid = str(p.get("uid") or "").strip()
        if uid:
            xk = uid
            note = "销客产品编码取自产品快照唯一性ID（金蝶物料表未覆盖该物料）"
        else:
            note = "销客产品编码待人工确认（金蝶物料表未覆盖、快照无唯一性ID）"
    return {"我方产品名称": p["name"], "我方物料名称": p["mat"], "物料编码": p["code"],
            "规格型号": p["spec"], "销客产品编码": xk, "所属组织": "",
            "匹配方式": method, "置信度": CONF[method],
            "来源": "产品主数据快照", "备注": note}


def _uniq(hits, method):
    """同名多条：销客编码相同视为同一条物料，否则算冲突。"""
    if len(hits) == 1:
        return _pack_mat(hits[0], method), []
    if len({h["xk"] for h in hits}) == 1:
        return _pack_mat(hits[0], method), []
    return None, hits


def match_one(row, materials, products, org=None, index=None):
    """返回 (结果 dict | None, 冲突候选 list, 被闸拦下的候选 list)。

    顺序固定（每级先查产品主数据快照、再查金蝶物料表）：
      1 物料名称精确 → 2 规格型号精确 → 3 型号标题模糊（候选必须过两道闸）。短路返回。
    """
    idx = index or build_index(materials, products, org)
    m_val = str(row.get("material") or "").strip()
    s_val = str(row.get("spec") or "").strip()
    model = str(row.get("model") or "").strip()

    # 1) 物料名称精确
    if m_val.upper() not in PLACEHOLDER:
        p = idx["prod_by_mat"].get(m_val) or idx["prod_by_name"].get(m_val)
        if p:
            return _pack_prod(p, "物料名称精确", idx), [], []
        res, conflict = _uniq(idx["by_name"].get(m_val, []), "物料名称精确")
        if res or conflict:
            return res, conflict, []

    # 2) 规格型号精确
    if s_val.upper() not in PLACEHOLDER:
        p = idx["prod_by_spec"].get(s_val)
        if p:
            return _pack_prod(p, "规格型号精确", idx), [], []
        res, conflict = _uniq(idx["by_spec"].get(s_val, []), "规格型号精确")
        if res or conflict:
            return res, conflict, []

    # 3) 型号标题模糊（仅型号标题可用时启用）
    if not model:
        return None, [], []
    if idx["has_products"]:
        name, _ = bi.match_product(model, products)
        if name:
            p = idx["prod_by_name"].get(name)
            if p:
                return _pack_prod(p, "型号标题模糊", idx), [], []
    blocked = []
    for spec, m in idx["spec_pairs"]:
        if spec not in model:
            continue
        ok, why = bi.fuzzy_ok(model, m["name"])
        if ok:
            return _pack_mat(m, "型号标题模糊"), [], []
        blocked.append({"候选": m["name"], "原因": why, "来源表": m["src"]})
    return None, [], blocked


def classify_miss(model, material_val, conflict, blocked):
    if conflict:
        return "存在同名多物料", "命中 %d 条同名：%s" % (
            len(conflict), "、".join(sorted({c["name"] for c in conflict})[:3]))
    if blocked:
        return "疑似错配", "候选 %s 被闸拦下（%s）" % (blocked[0]["候选"], blocked[0]["原因"])
    return "我方无此物料", "物料表与产品快照均无命中；判定前请过连接器确认"


# ---------------------------------------------------------------- 输出
def write_report(path, ext_hdr, ext_rows, results, stats):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="2F5597")

    ws = wb.active
    ws.title = "匹配结果"
    ws.append([str(c) if c is not None else "" for c in ext_hdr] + ADD_COLS)
    for raw, res in zip(ext_rows, results):
        ws.append([("" if v is None else v) for v in raw] + [res.get(k, "") for k in ADD_COLS])
    for i in range(1, len(ext_hdr) + len(ADD_COLS) + 1):
        ws.cell(1, i).font = head
        ws.cell(1, i).fill = fill
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("未匹配清单")
    ws2.append(["原因分类", "外部型号标题", "外部物料名称", "说明"])
    for r in stats["未匹配明细"]:
        ws2.append([r["原因"], r["model"], r["material"], r["说明"]])
    for i in range(1, 5):
        ws2.cell(1, i).font = head
        ws2.cell(1, i).fill = fill
    for col, w in zip("ABCD", (22, 40, 44, 46)):
        ws2.column_dimensions[col].width = w

    ws3 = wb.create_sheet("统计")
    ws3.append(["项目", "数值"])
    for k, v in stats["概况"].items():
        ws3.append([k, v])
    marks = []
    for title, dist in (("匹配方式", stats["方式分布"]), ("命中来源", stats["来源分布"]),
                        ("未匹配原因", stats["原因分布"])):
        ws3.append(["", ""])
        marks.append(ws3.max_row + 1)
        ws3.append([title, "行数"])
        for k, v in dist.items():
            ws3.append([k, v])
    for r in [1] + marks:
        for c in (1, 2):
            ws3.cell(r, c).font = head
            ws3.cell(r, c).fill = fill
    ws3.column_dimensions["A"].width = 34
    ws3.column_dimensions["B"].width = 18
    wb.save(path)


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--source", required=True, help="外部物料表")
    ap.add_argument("--sheet", default=None, help="不传则按表头语义自动选择")
    ap.add_argument("--org", default=None, help="按组织过滤我方物料表（不传则跨组织匹配）")
    ap.add_argument("--material-table", nargs="*", default=None, help="金蝶物料表，可多张")
    ap.add_argument("--product-snapshot", default=None)
    ap.add_argument("--downloads", default=r"D:\Backup\Downloads")
    ap.add_argument("--outdir", default=None)
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    if not os.path.exists(args.source):
        print(json.dumps({"error": "外部物料表不存在: %s" % args.source}, ensure_ascii=False))
        sys.exit(2)

    mt_paths = args.material_table or default_material_tables(args.downloads)
    if isinstance(mt_paths, str):
        mt_paths = [mt_paths]
    mt_paths = [p for p in mt_paths if os.path.exists(p)]
    materials = []
    for p in mt_paths:
        materials.extend(load_material_table(p))

    snap = args.product_snapshot
    if not snap:
        cands = sorted(glob.glob(os.path.join(args.downloads, "产品对象导出结果_*.xlsx")), reverse=True)
        snap = next((f for f in cands if not os.path.basename(f).startswith("~$")), None)
    has_snap = bool(snap and os.path.exists(snap))
    products = bi.load_products(snap) if has_snap else []

    # 金蝶物料表缺失 → 降级（只匹配产品快照，编码回落快照唯一性ID）；
    # 两者都没有才停手——一个底表都没有时匹配结果没有依据。
    degrade = ""
    if not materials:
        if not has_snap:
            print(json.dumps({"error": "既没有可用的金蝶物料表，也没有产品快照，本能力不可用"
                                       "（不做无依据的猜测匹配）"}, ensure_ascii=False))
            sys.exit(2)
        degrade = ("金蝶物料表缺失：降级为只匹配产品快照，销客产品编码取自快照唯一性ID，"
                   "物料编码取自快照物料编号；结果仅供参考，导入前请人工复核")

    wb = bi.open_wb(args.source)
    sheet_name = args.sheet or pick_sheet(wb)
    if not sheet_name or sheet_name not in wb.sheetnames:
        print(json.dumps({"error": "定位不到外部表的数据 sheet（现有: %s），请用 --sheet 指定"
                          % wb.sheetnames}, ensure_ascii=False))
        sys.exit(2)
    ws = wb[sheet_name]
    rows = read_rows(ws, 1)
    hdr_row = find_header_row(rows)
    hdr = row_at(rows, hdr_row)
    cols = resolve_columns(hdr)
    # 裁掉空尾巴列：有的源表被 Excel 记成 16384 列，全写进输出会让 297 行的报告变成 13 MB。
    last = _last_content_col(rows, hdr_row)
    if last and last < len(hdr):
        hdr = hdr[:last]
        rows = [(i, v[:last]) for i, v in rows]
    if not any(cols.get(k) for k in KEY_COLS):
        print(json.dumps({"error": "外部表里型号类/物料名称类/规格型号类列一个都没找到，"
                                   "无法匹配；表头行=%d，表头=%s"
                          % (hdr_row, [str(x) for x in hdr[:12]])}, ensure_ascii=False))
        sys.exit(2)

    matcher = Matcher(materials, products)
    index = matcher.index(args.org)

    def val(vals, key):
        ci = cols.get(key)
        v = vals[ci - 1] if ci and ci <= len(vals) else None
        return "" if v is None else str(v).strip()

    ext_hdr = [c for c in hdr]
    ext_rows, results, miss_rows = [], [], []
    n_skip = 0
    for ridx, vals in rows:
        if ridx <= hdr_row:
            continue
        model = val(vals, "model")
        mval = val(vals, "material")
        if not model and mval.upper() in PLACEHOLDER:
            n_skip += 1
            continue
        row = {"model": model, "material": mval, "spec": val(vals, "spec"),
               "brand": val(vals, "brand"), "industry": val(vals, "industry")}
        res, conflict, blocked = match_one(row, materials, products, args.org, index)
        if res is None:
            reason, detail = classify_miss(model, mval, conflict, blocked)
            res = {"匹配方式": "未匹配", "置信度": 0.0}
            miss_rows.append({"原因": reason, "model": model, "material": mval, "说明": detail})
        ext_rows.append(vals + [None] * max(0, len(ext_hdr) - len(vals)))
        results.append(res)

    from collections import Counter
    n_matched = sum(1 for r in results if r["匹配方式"] != "未匹配")
    n_no_code = sum(1 for r in results if r["匹配方式"] != "未匹配" and not r.get("销客产品编码"))
    stats = {
        "概况": {
            "外部表": os.path.basename(args.source), "sheet": sheet_name,
            "表头行": hdr_row, "数据行": len(results), "跳过空行/占位行": n_skip,
            "组织过滤": args.org or "未指定（跨组织）",
            "我方物料表": "、".join(os.path.basename(p) for p in mt_paths),
            "我方物料条数": len(materials), "参与匹配条数": len(index["pool"]),
            "产品快照": os.path.basename(snap) if snap else "无（第三级模糊仅走物料表）",
            "匹配上": n_matched, "未匹配": len(miss_rows),
            "已匹配但取不到销客产品编码": n_no_code,
        },
        "方式分布": dict(Counter(r["匹配方式"] for r in results).most_common()),
        "来源分布": dict(Counter(r.get("来源", "") for r in results).most_common()),
        "原因分布": dict(Counter(m["原因"] for m in miss_rows).most_common()),
        "未匹配明细": miss_rows,
    }
    if degrade:
        stats["概况"]["降级说明"] = degrade

    outdir = args.outdir or os.path.dirname(os.path.abspath(args.source))
    os.makedirs(outdir, exist_ok=True)
    if args.dry_run:
        files = ["(dry-run 未写文件)"]
    else:
        p = os.path.join(outdir, "外部物料匹配结果_%s_%s.xlsx"
                         % (args.org or "跨组织", datetime.date.today().strftime("%Y%m%d")))
        write_report(p, ext_hdr, ext_rows, results, stats)
        files = [p]
    wb.close()

    out = dict(stats["概况"])
    out.update({"匹配方式分布": stats["方式分布"], "未匹配原因分布": stats["原因分布"],
                "列位置": cols, "files": files})
    print(json.dumps(out, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
