#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""build_import.py 的行为测试（TDD）。

运行： python scripts/test_build_import.py
原则：先让这些用例变绿，才允许改 build_import.py 的实现。

覆盖两个核心口径：
  A. 组织(org) 决定「价目表」列；品牌 是产品实际品牌，与组织无关。
     - 欣暖家组织 → 美的价目表20250901；产品 TCL 彩电 → 品牌列 TCL
     - 乾鑫组织   → 乾鑫价目表20260701；产品 美的空调 → 品牌列 美的
     - 西门子组织 → 西门子价目表20260701；产品 西门子洗碗机 → 品牌列 西门子
  B. 品牌词表需覆盖 COLMO/东芝/索尼/松下/卡萨帝/科沃斯/长虹/容声/史密斯 等。
"""
import os
import sys
import json
import shutil
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_import as bi  # noqa: E402


class TestBrandInference(unittest.TestCase):
    """品牌 = 产品实际品牌，从产品名/型号推断。"""

    def test_common_brands(self):
        cases = [
            ("TCL平板电视98C10L-A", "TCL"),
            ("美的3匹变频柜机KFR-72LW/FQ1-1", "美的"),
            ("小天鹅滚筒洗干一体机TD10V628PLUS", "小天鹅"),
            ("西门子嵌入式洗碗机SJ43EB24KC", "西门子"),
            ("格力内机GMV-NHDR80PS/Fb无电辅带水泵", "格力"),
            ("方太油烟机CXW-358-03-X1A-W", "方太"),
            ("大金空调RJSYQ140AAV", "大金"),
        ]
        for model, expect in cases:
            with self.subTest(model=model):
                self.assertEqual(bi.guess_brand_from_model(model), expect)

    def test_newly_added_brands(self):
        """这些品牌曾因不在词表里而被错误回退成组织名，必须能推断出来。"""
        cases = [
            ("COLMO多开门冰箱CRBUF706N-X1熔幔岩", "COLMO"),
            ("东芝主机XCY-MGP2241HTN-C", "东芝"),
            ("东芝内机XMD-GP0271MHN-C", "东芝"),
            ("东芝线控器RBC-ASXG31N-C", "东芝"),
            ("索尼电视机98XR70M2", "索尼"),
            ("松下多门冰箱NR-F611GA1-H", "松下"),
            ("卡萨帝壁挂炉LL1PBD26-CECONPXGU1", "卡萨帝"),
            ("科沃斯T90PRO水箱板白色", "科沃斯"),
            ("长虹电视机75Q70S", "长虹"),
            ("容声多门冰箱BCD-505Q50CZLBD", "容声"),
            ("史密斯净水机WF30D1", "史密斯"),
        ]
        for model, expect in cases:
            with self.subTest(model=model):
                self.assertEqual(bi.guess_brand_from_model(model), expect,
                                 f"品牌词表缺少 {expect}：{model}")

    def test_tmall_inventory_20260928_brands(self):
        """2026-09-28 天猫库存物料价目表实测未识别的 104 行，涉及以下品牌。"""
        cases = [
            ("樱花13升燃气热水器SCH-13Q125A天燃气", "樱花"),
            ("美菱三开门冰箱MRF-226WP3CXG1", "美菱"),
            ("火星人集成灶ET70BC天然气", "火星人"),
            ("奥普浴霸A127", "奥普"),
            ("瑞尔特浴室柜组合Y1009玄铁灰", "瑞尔特"),
            ("万家乐两眼灶JZT-V7-12T", "万家乐"),
            ("易开得台上式净水器 C1 Pro", "易开得"),
            ("三星多开门冰箱RF50DG5151QQSC石岩灰", "三星"),
            ("九阳0涂层电饭煲40N5U", "九阳"),
            ("凯迪仕智能锁A5-W", "凯迪仕"),
            ("博乐宝中央净水器Z10灰色", "博乐宝"),
            ("奥克斯1.5P挂机KFR-35GW/BpR3AEH29(B1)", "奥克斯"),
            ("沁园反渗透净水器KRL3933灰色", "沁园"),
            ("云米洗干套装Master 2 MAX", "云米"),
            ("华生电风扇   5V-8", "华生"),
            ("康宝消毒柜XDZ168-SPACEX", "康宝"),
            ("百得欧式油烟机CXW-250-AC13Pro", "百得"),
            ("统帅洗烘一体机XQGL125-MBLDE697WU1", "统帅"),
            ("苏泊尔电烤箱OJ38A02", "苏泊尔"),
            ("夸克 AI眼镜", "夸克"),
            ("xRealVR头盔Beam Pro", "xReal"),
            ("弗迪沃斯除湿新风FD-X128P", "弗迪沃斯"),
        ]
        for model, expect in cases:
            with self.subTest(model=model):
                self.assertEqual(bi.guess_brand_from_model(model), expect,
                                 f"品牌词表缺少 {expect}：{model}")

    def test_org_name_is_never_used_as_brand(self):
        """硬规则：组织名（桦润/天猫/芯猫）永远不能当品牌推断结果。"""
        self.assertNotIn("桦润", bi.BRAND_KEYWORDS)
        self.assertNotIn("天猫", bi.BRAND_KEYWORDS)
        self.assertNotIn("芯猫", bi.BRAND_KEYWORDS)

    def test_unknown_brand_returns_none(self):
        self.assertIsNone(bi.guess_brand_from_model("某某没听过的型号XYZ"))


class TestClassify(unittest.TestCase):
    """分类特例：乾鑫读源表、大金固定空调、欣暖家留空。"""

    def test_xinnuanjia_org_leaves_classification_blank(self):
        self.assertEqual(bi.classify("TCL平板电视98C10L-A", "欣暖家", "", ""),
                         ("", "", True))

    def test_dajin_fixed_aircon(self):
        self.assertEqual(bi.classify("大金某机型", "大金", "", ""),
                         ("空调", "空调", True))

    def test_qianxin_reads_source_columns(self):
        self.assertEqual(
            bi.classify("小天鹅1.5匹变频挂机", "乾鑫", "空调", "柜挂天井风管机家用多联"),
            ("空调", "柜挂天井风管机家用多联", True))

    def test_siemens_cross_fill(self):
        """西门子：大类列填【西门子品项】组，品项列填【西门子大类】组。"""
        pclass, pitem, matched = bi.classify("西门子微蒸烤一体机CP1K4R7T7W", "西门子", "", "")
        self.assertTrue(matched)
        self.assertEqual(pclass, "微蒸烤【西门子品项】")
        self.assertEqual(pitem, "微蒸烤箱【西门子大类】")


class TestRoundPrice(unittest.TestCase):
    def test_round_half_up(self):
        self.assertEqual(bi.round_price(423.52), 424)
        self.assertEqual(bi.round_price(2214.53), 2215)
        self.assertEqual(bi.round_price(14267.18), 14267)
        self.assertEqual(bi.round_price(6658.88888888889), 6659)
        self.assertEqual(bi.round_price(1.2), 1)
        self.assertEqual(bi.round_price(1.6), 2)

    def test_none_passthrough(self):
        self.assertIsNone(bi.round_price(None))


class TestGrade(unittest.TestCase):
    def test_single_letter(self):
        self.assertEqual(bi.extract_grade("D"), "D类")
        self.assertEqual(bi.extract_grade("a"), "A类")

    def test_empty_note(self):
        self.assertEqual(bi.extract_grade(None), "")
        self.assertEqual(bi.extract_grade(""), "")
        self.assertEqual(bi.extract_grade("2025-03-26 00:00:00"), "")


class TestProductMatching(unittest.TestCase):
    """产品匹配：不得把 F60-33Q7Pro(HE) 错配到 F60-33Q7Pro。"""

    PRODUCTS = [
        {"name": "美的电热水器F60-33Q7Pro", "mat": "美的电热水器F60-33Q7Pro",
         "spec": "F60-33Q7Pro", "code": "20.04.01.01.079",
         "pclass": "", "pitem": ""},
        {"name": "美的电热水器F60-33Q7Pro(HE)", "mat": "美的电热水器F60-33Q7Pro(HE)",
         "spec": "F60-33Q7Pro(HE)", "code": "20.04.01.01.099",
         "pclass": "", "pitem": ""},
    ]

    def test_he_suffix_not_mismatched(self):
        """源表要 (HE) 版本，绝不能返回不带 (HE) 的那条。"""
        name, how = bi.match_product("美的电热水器F60-33Q7Pro(HE)", self.PRODUCTS)
        self.assertEqual(name, "美的电热水器F60-33Q7Pro(HE)",
                         "错配到了不带 (HE) 的产品")

    def test_plain_version_still_matches(self):
        name, how = bi.match_product("美的电热水器F60-33Q7Pro", self.PRODUCTS)
        self.assertEqual(name, "美的电热水器F60-33Q7Pro")

    # ---- 2026-09-15 实战错配回归 ----

    def test_brand_mismatch_blocked(self):
        """东芝…100Z600QFR-S 绝不能命中 索尼…100Z600QF（品牌闸）。"""
        prods = [{"name": "索尼平板电视100Z600QF", "mat": "索尼平板电视100Z600QF",
                  "spec": "100Z600QF", "code": "", "pclass": "", "pitem": ""}]
        name, _ = bi.match_product("东芝平板电视100Z600QFR-S", prods)
        self.assertIsNone(name)

    def test_generic_name_blocked(self):
        """候选名无型号码（如「波轮洗衣机」）不得作为命中。"""
        prods = [{"name": "波轮洗衣机", "mat": "波轮洗衣机",
                  "spec": "", "code": "", "pclass": "", "pitem": ""}]
        name, _ = bi.match_product("美的波轮洗衣机MB10V56T", prods)
        self.assertIsNone(name)

    def test_prefix_sibling_blocked(self):
        """MB- 前缀并存产品：源表 MB-AFB40C8 不能命中并存的 AFB40C8（边界闸）。"""
        prods = [{"name": "美的电饭煲AFB40C8", "mat": "美的电饭煲AFB40C8",
                  "spec": "AFB40C8", "code": "", "pclass": "", "pitem": ""}]
        name, _ = bi.match_product("美的电饭煲MB-AFB40C8", prods)
        self.assertIsNone(name)

    def test_roman_suffix_blocked(self):
        """型号码后紧跟罗马数字（Ⅱ vs 无后缀）不算命中。"""
        prods = [{"name": "美的空调MJV-200W-E01-LH", "mat": "美的空调MJV-200W-E01-LH",
                  "spec": "MJV-200W-E01-LH", "code": "", "pclass": "", "pitem": ""}]
        name, _ = bi.match_product("美的空调MJV-200W-E01-LHⅡ", prods)
        self.assertIsNone(name)

    def test_fuzzy_still_hits_when_clean(self):
        """闸不能误伤：空格差异的干净命中要照常返回。"""
        prods = [{"name": "美的电饭煲MB-AFB40C8", "mat": "美的电饭煲MB-AFB40C8",
                  "spec": "MB-AFB40C8", "code": "", "pclass": "", "pitem": ""}]
        name, how = bi.match_product("美的电饭煲 MB-AFB40C8", prods)
        self.assertEqual(name, "美的电饭煲MB-AFB40C8")


class TestEndToEnd(unittest.TestCase):
    """端到端：组织决定价目表，品牌列是产品实际品牌。"""

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="plit_test_")
        self.src = os.path.join(self.tmp, "源表.xlsx")
        self.prod = os.path.join(self.tmp, "产品快照.xlsx")
        self.det = os.path.join(self.tmp, "明细快照.xlsx")
        self.outdir = os.path.join(self.tmp, "out")
        self._make_source()
        self._make_product_snapshot()
        self._make_detail_snapshot()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _make_source(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "申请表"
        ws.append(["纷享销客价目表新增（变更）申请表"])
        ws.append(["申请部门：欣暖家"])
        ws.append(["存货编码", "产品官方名称及型号", "零售建议价", "零售限制价",
                   "销售结算价", "实际成本价", "备注"])
        ws.append(["", "TCL平板电视98C10L-A", 21999, 14549, 12800, 12800, ""])
        ws.append(["", "COLMO多开门冰箱CRBUF706N-X1熔幔岩", 14400, 10666, 9600, 9600, ""])
        wb.save(self.src)

    def _make_product_snapshot(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "产品数据"
        ws.append(["产品名称（必填）", "物料名称", "型号规格", "物料编号",
                   "产品大类", "产品品相"])
        for n in ["TCL平板电视98C10L-A", "COLMO多开门冰箱CRBUF706N-X1熔幔岩"]:
            ws.append([n, n, n, "X", "", ""])
        wb.save(self.prod)

    def _make_detail_snapshot(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "价目表明细数据"
        ws.append(["唯一性ID（必填）", "价目表明细编号", "产品（必填）", "价目表（必填）",
                   "价目表_唯一性ID（必填）"])
        wb.save(self.det)

    def _run(self, org):
        import subprocess
        cmd = [sys.executable, os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                            "build_import.py"),
               "--source", self.src, "--org", org,
               "--product-snapshot", self.prod, "--detail-snapshot", self.det,
               "--outdir", self.outdir]
        r = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8")
        return r

    def _read_out(self):
        from openpyxl import load_workbook
        import glob
        fs = glob.glob(os.path.join(self.outdir, "*新建*.xlsx"))
        self.assertTrue(fs, "未生成新建文件")
        wb = load_workbook(fs[0], data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        hdr = [str(c) if c is not None else "" for c in rows[0]]
        return hdr, [dict(zip(hdr, r)) for r in rows[1:]]

    def test_xinnuanjia_org_uses_midea_pricebook_and_tcl_brand(self):
        r = self._run("欣暖家")
        self.assertEqual(r.returncode, 0, r.stderr)
        hdr, data = self._read_out()
        self.assertEqual(len(data), 2)
        by_prod = {d["产品（必填）"]: d for d in data}
        self.assertEqual(by_prod["TCL平板电视98C10L-A"]["品牌"], "TCL")
        self.assertEqual(by_prod["TCL平板电视98C10L-A"]["价目表（必填）"],
                         "美的价目表20250901")
        self.assertEqual(by_prod["COLMO多开门冰箱CRBUF706N-X1熔幔岩"]["品牌"], "COLMO")
        # 欣暖家组织 → 大类/品项留空
        self.assertIn(by_prod["TCL平板电视98C10L-A"]["产品大类"], (None, ""))
        self.assertIn(by_prod["TCL平板电视98C10L-A"]["产品品项"], (None, ""))

    def test_fixed_columns(self):
        self._run("欣暖家")
        hdr, data = self._read_out()
        for d in data:
            self.assertEqual(d["价目表折扣（%）"], 100)
            self.assertEqual(d["业务类型（必填）"], "预设业务类型")
            self.assertEqual(d["价目表售价"], 21999 if "TCL" in d["产品（必填）"] else 14400)

    def test_siemens_org_org_pricebook(self):
        """西门子组织 → 西门子价目表20260701。"""
        r = self._run("西门子")
        self.assertEqual(r.returncode, 0, r.stderr)
        hdr, data = self._read_out()
        for d in data:
            self.assertEqual(d["价目表（必填）"], "西门子价目表20260701")


class TestSemanticHeader(unittest.TestCase):
    """表头按语义识别：不写死格式名、不写死行号与列字母。"""

    @staticmethod
    def _wb(rows, sheet="数据", extra=None):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = sheet
        for r in rows:
            ws.append(r)
        for name, rws in (extra or []):
            w = wb.create_sheet(name)
            for r in rws:
                w.append(r)
        return wb

    def test_norm_header_strips_parentheses(self):
        cases = [("产品（必填）", "产品"), ("规格型号（单列）", "规格型号"),
                 ("（浮动下限价）品牌负责人审核价", "品牌负责人审核价"),
                 ("价目表折扣（%）", "价目表折扣"),
                 ("物料名称（芯猫+彩电物料表）", "物料名称")]
        for raw, expect in cases:
            with self.subTest(raw=raw):
                self.assertEqual(bi.norm_header(raw), expect)

    def test_sales_limit_price_column_is_limit(self):
        """「销售限价」要认成 limit（= 品牌负责人审核价）。

        2026-09-28 天猫库存物料价目表用的就是这个列名，此前只认「销售限制价」、
        而「限价」是 2 字短词只认整格相等 → 整列漏掉，导入文件限价 379 行全空。
        """
        for col in ["销售限价", "销售限制价", "零售限制价", "品牌负责人审核价"]:
            with self.subTest(col=col):
                hdr = ["物料编码", "物料名称", "成本价", "销售结算价", col, "零售指导价"]
                self.assertEqual(bi.col_index(hdr, "limit"), 4, f"{col} 未被认成限价列")

    def test_short_keyword_needs_exact_match(self):
        """「产品」不能命中「产品大类」，否则表头行会数错列、列位会串。"""
        self.assertFalse(bi.header_hit("产品大类", "产品"))
        self.assertFalse(bi.header_hit("产品品项", "产品"))
        self.assertTrue(bi.header_hit("产品（必填）", "产品"))
        self.assertTrue(bi.header_hit("产品大类", "产品大类"))

    def test_row2_header_detected(self):
        """R1 是标题行或空行时，表头在 R2、数据 R3 起。"""
        wb = self._wb([[None, None, None],
                       ["品牌", "型号标题", "销售指导价"],
                       ["Ronshen/容声", "BCD-515W60FZBAS琥珀钰", 3999]])
        hdr_row, data_row, hdr = bi.find_header_row(wb["数据"])
        self.assertEqual((hdr_row, data_row), (2, 3))
        self.assertEqual(bi.col_index(hdr, "model"), 1)
        self.assertEqual(bi.col_index(hdr, "price"), 2)
        self.assertEqual(bi.col_index(hdr, "brand"), 0)

    def test_apply_form_header_at_r3(self):
        wb = self._wb([["纷享销客价目表新增（变更）申请表"],
                       ["申请部门：欣暖家"],
                       ["存货编码", "产品官方名称及型号", "零售建议价"],
                       ["", "TCL平板电视98C10L-A", 21999]])
        fmt, hdr_row, data_row, hdr = bi.detect_format_and_rows(wb["数据"])
        self.assertEqual(fmt, "apply")
        self.assertEqual((hdr_row, data_row), (3, 4))
        self.assertEqual(bi.detect_org(wb["数据"], hdr_row), "欣暖家")

    def test_template_header_at_r1(self):
        wb = self._wb([["唯一性ID（必填）", "产品（必填）", "价目表售价", "价目表（必填）"],
                       ["x", "TCL平板电视98C10L-A", 21999, "美的价目表20250901"]])
        fmt, hdr_row, data_row, hdr = bi.detect_format_and_rows(wb["数据"])
        self.assertEqual((fmt, hdr_row, data_row), ("template", 1, 2))
        self.assertEqual(bi.col_index(hdr, "model"), 1)

    def test_r1_empty_sheet_selected_and_sheet2_ignored(self):
        """R1 全空但 R2 是表头 → 要选中真数据表；Sheet2 这类空壳表不能被选中。"""
        wb = self._wb([[None, None, None],
                       ["品牌", "型号标题", "销售指导价"],
                       ["容声", "BCD-606WKK1FPGZA", 5999]],
                      extra=[("Sheet2", [["随便写点草稿"]]), ("Sheet1", [["草稿"]])])
        ws, name = bi.find_data_sheet(wb)
        self.assertEqual(name, "数据")

    def test_header_failure_raises_without_guessing(self):
        wb = self._wb([["甲", "乙"], ["1", "2"]])
        with self.assertRaises(ValueError):
            bi.find_header_row(wb["数据"])

    def test_only_sheet_named_sheet1_with_real_data_is_used(self):
        """整个文件只有一个 Sheet1 且装着真数据 → 必须用它，不能因名字跳过。
        （实测「泽锋新增 (17).xlsx」就是这种：唯一 Sheet1，品牌/物料名称/规格型号/价。）
        """
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["品牌", "物料名称", "规格型号", "供 价(RMB)", "核算价", "限价", "面价", "品类"])
        ws.append(["美的", "美的内机MDV-D22Q4/BP", "MDV-D22Q4/BP3N1-", 1670, 1670, None, 3340, "空调"])
        got, name = bi.find_data_sheet(wb)
        self.assertEqual(name, "Sheet1")
        self.assertIsNotNone(got)
        fmt, hdr_row, data_row, hdr = bi.detect_format_and_rows(got)
        self.assertEqual(fmt, "generic")
        self.assertEqual(hdr_row, 1)
        # col_index 返回 0 基下标（与 iter_rows 元组一致，见 load_source_rows 的 r[i] 用法）
        self.assertEqual(bi.col_index(hdr, "model"), 2)
        self.assertEqual(bi.col_index(hdr, "price"), 6, "面价 → 价目表售价")
        self.assertEqual(bi.col_index(hdr, "cost"), 3, "供 价(RMB) → 成本价")
        self.assertEqual(bi.col_index(hdr, "settle"), 4, "核算价 → 结算价")
        self.assertEqual(bi.col_index(hdr, "limit"), 5, "限价 → 限价")

    def test_named_sheet_beats_sheet1_with_semantics(self):
        """真数据表与带语义的 Sheet1 复制品并存 → 优先用非 SheetN 名的。"""
        wb = self._wb([["品牌", "型号标题", "销售指导价"],
                       ["容声", "BCD-606WKK1FPGZA", 5999]])
        cp = wb.create_sheet("Sheet1")
        cp.append(["品牌", "型号标题", "销售指导价"])
        cp.append(["草稿复制品", "X", 1])
        ws, name = bi.find_data_sheet(wb)
        self.assertEqual(name, "数据")

    def test_rightmost_dated_sheet_wins(self):
        """多 sheet 累积表：只取最右侧日期 sheet，其余不合并。"""
        wb = self._wb([["产品官方名称及型号", "零售建议价"], ["A", 1]], sheet="260901")
        for name in ["260910", "260915"]:
            w = wb.create_sheet(name)
            w.append(["产品官方名称及型号", "零售建议价"])
            w.append(["A", 1])
        ws, name = bi.find_data_sheet(wb)
        self.assertEqual(name, "260915")


class TestBrandResolution(unittest.TestCase):
    """品牌 = 产品实际品牌；推不出来就留空，绝不回退成组织名。"""

    def test_slash_form_takes_chinese_part(self):
        cases = [("Ronshen/容声", "BCD-606WKK1FPGZA", "容声"),
                 ("Casarte/卡萨帝", "BCD-633WLCFDANZCU1", "卡萨帝"),
                 ("SIEMENS/西门子", "KF88HVA56C", "西门子")]
        for raw, model, expect in cases:
            with self.subTest(raw=raw):
                self.assertEqual(bi.resolve_brand(raw, model), expect)

    def test_no_chinese_name_used_as_is(self):
        self.assertEqual(bi.resolve_brand("TCL", "TCL平板电视98C10L-A"), "TCL")

    def test_english_only_falls_back_to_model(self):
        """纯英文品牌列（Siemens）先用型号推中文名。"""
        self.assertEqual(bi.resolve_brand("Siemens", "西门子嵌入式洗碗机SJ43EB24KC"), "西门子")

    def test_never_falls_back_to_org(self):
        """推不出品牌必须返回空串，不能返回组织名。"""
        self.assertEqual(bi.resolve_brand("", "某不知名热水器X100"), "")
        self.assertEqual(bi.resolve_brand("", ""), "")

    def test_source_brand_column_wins(self):
        self.assertEqual(bi.resolve_brand("海尔", "容声多门冰箱BCD-505"), "海尔")


class TestTmallXinmaoOrg(unittest.TestCase):
    """天猫与芯猫是同一个组织，映射到同一张价目表。"""

    def test_both_aliases_same_pricebook(self):
        self.assertEqual(bi.ORG_PRICEBOOKS["天猫"], "天猫价目表20260901")
        self.assertEqual(bi.ORG_PRICEBOOKS["芯猫"], "天猫价目表20260901")


class TestGenericFormatEndToEnd(unittest.TestCase):
    """货品表（R1 空、表头 R2、多 sheet）+ 天猫组织：价目表取天猫，品牌取产品实际品牌。"""

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="plit_gen_")
        self.src = os.path.join(self.tmp, "货品表.xlsx")
        self.prod = os.path.join(self.tmp, "产品快照.xlsx")
        self.det = os.path.join(self.tmp, "明细快照.xlsx")
        self.outdir = os.path.join(self.tmp, "out")
        self._make_source()
        self._make_product_snapshot()
        self._make_detail_snapshot()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _make_source(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "9月货品表"
        ws.append([None, None, None, None, None, None, None])
        ws.append(["品牌", "行业", "型号标题", "销售指导价", "销售结算价", "限价", "云仓底价"])
        ws.append(["Ronshen/容声", "冰箱", "容声多开门冰箱BCD-606WKK1FPGZA", 5999, 4200, 5600, 4100])
        ws.append(["TCL", "电视机", "TCL平板电视98C10L-A", 21999, 12800, 14549, 12800])
        wb.create_sheet("Sheet2").append(["草稿页，不是数据"])
        wb.save(self.src)

    def _make_product_snapshot(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "产品数据"
        ws.append(["产品名称（必填）", "物料名称", "型号规格", "物料编号", "产品大类", "产品品相"])
        for n in ["容声多开门冰箱BCD-606WKK1FPGZA", "TCL平板电视98C10L-A"]:
            ws.append([n, n, n, "X", "", ""])
        wb.save(self.prod)

    def _make_detail_snapshot(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "价目表明细数据"
        ws.append(["唯一性ID（必填）", "价目表明细编号", "产品（必填）", "价目表（必填）",
                   "价目表_唯一性ID（必填）"])
        wb.save(self.det)

    def _run(self, org):
        import subprocess
        cmd = [sys.executable, os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                            "build_import.py"),
               "--source", self.src, "--org", org,
               "--product-snapshot", self.prod, "--detail-snapshot", self.det,
               "--outdir", self.outdir]
        return subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8")

    def _read_out(self):
        from openpyxl import load_workbook
        import glob
        fs = glob.glob(os.path.join(self.outdir, "*新建*.xlsx"))
        self.assertTrue(fs, "未生成新建文件")
        wb = load_workbook(fs[0], data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        hdr = [str(c) if c is not None else "" for c in rows[0]]
        return hdr, [dict(zip(hdr, r)) for r in rows[1:]]

    def test_tmall_org_pricebook_and_real_brands(self):
        r = self._run("天猫")
        self.assertEqual(r.returncode, 0, r.stderr)
        hdr, data = self._read_out()
        self.assertEqual(len(data), 2, "货品表两行都要生成")
        by_prod = {d["产品（必填）"]: d for d in data}
        self.assertIn("容声多开门冰箱BCD-606WKK1FPGZA", by_prod)
        self.assertIn("TCL平板电视98C10L-A", by_prod)
        for d in data:
            self.assertEqual(d["价目表（必填）"], "天猫价目表20260901")
        # 品牌 = 产品实际品牌（「Ronshen/容声」取中文段），不是组织名
        self.assertEqual(by_prod["容声多开门冰箱BCD-606WKK1FPGZA"]["品牌"], "容声")
        self.assertEqual(by_prod["TCL平板电视98C10L-A"]["品牌"], "TCL")

    def test_xinmao_alias_same_result(self):
        r = self._run("芯猫")
        self.assertEqual(r.returncode, 0, r.stderr)
        hdr, data = self._read_out()
        for d in data:
            self.assertEqual(d["价目表（必填）"], "天猫价目表20260901")

    def test_json_summary_reports_header_row(self):
        r = self._run("天猫")
        payload = json.loads(r.stdout[r.stdout.index("{"):])
        self.assertEqual(payload["format"], "generic")
        self.assertEqual(payload["sheet"], "9月货品表")
        self.assertEqual(payload["新建"], 2)


if __name__ == "__main__":
    unittest.main(verbosity=2)
