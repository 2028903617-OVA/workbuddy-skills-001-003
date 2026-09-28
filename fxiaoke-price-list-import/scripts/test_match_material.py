#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""match_material.py 的行为测试（TDD）。

运行： python scripts/test_match_material.py
原则：先让这些用例变绿，才允许改 match_material.py 的实现。

覆盖的核心口径：
  1. 匹配顺序固定：物料名称精确 → 规格型号精确 → 型号标题模糊（短路）
  2. 型号标题不能当主键：只有型号标题时走模糊，且必须过两道闸
  3. 占位值（0/空）不是物料名，不得拿去匹配
  4. 同名多物料不许自动取第一条（销客编码相同才视为同一条）
  5. 列按表头语义定位，不按列字母
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import match_material as mm  # noqa: E402


def mat(name, spec="", xk="", org="桦润芯猫", code="01.01.01.01.02.001"):
    return {"src": "物料表.xlsx", "org": org, "code": code, "name": name,
            "spec": spec, "xk": xk, "group": ""}


class TestMatchOrder(unittest.TestCase):
    """匹配顺序：物料名称精确 → 规格型号精确 → 型号标题模糊。"""

    def setUp(self):
        self.materials = [
            mat("容声多开门冰箱BCD-606WKK1FPGZA", "BCD-606WKK1FPGZA", "6a41aa01", code="A1"),
            mat("TCL平板电视98C10L-A", "98C10L-A", "6a41aa02", code="A2"),
            mat("海信平板电视75D60S", "75D60S", "6a41aa03", code="A3"),
        ]

    def test_material_name_exact_wins(self):
        row = {"model": "BCD-606WKK1FPGZA曜金印", "material": "容声多开门冰箱BCD-606WKK1FPGZA"}
        res, conflict, blocked = mm.match_one(row, self.materials, [])
        self.assertEqual(res["匹配方式"], "物料名称精确")
        self.assertEqual(res["销客产品编码"], "6a41aa01")
        self.assertEqual(res["物料编码"], "A1")
        self.assertEqual((conflict, blocked), ([], []))

    def test_spec_exact_when_material_name_absent(self):
        row = {"model": "TCL平板电视98C10L-A", "material": "", "spec": "98C10L-A"}
        res, _, _ = mm.match_one(row, self.materials, [])
        self.assertEqual(res["匹配方式"], "规格型号精确")
        self.assertEqual(res["销客产品编码"], "6a41aa02")

    def test_material_name_beats_spec(self):
        """两列都能命中时，物料名称优先（命中率更高、更可信）。"""
        row = {"model": "TCL平板电视98C10L-A", "material": "TCL平板电视98C10L-A",
               "spec": "98C10L-A"}
        res, _, _ = mm.match_one(row, self.materials, [])
        self.assertEqual(res["匹配方式"], "物料名称精确")

    def test_placeholder_zero_is_not_a_material_name(self):
        """外部表的 0 是「我方没有」的占位，跳过它走下一级。"""
        row = {"model": "TCL平板电视98C10L-A", "material": "0", "spec": "98C10L-A"}
        res, _, _ = mm.match_one(row, self.materials, [])
        self.assertEqual(res["匹配方式"], "规格型号精确")

    def test_unknown_model_not_matched(self):
        row = {"model": "某不知名热水器X999", "material": "", "spec": ""}
        res, conflict, blocked = mm.match_one(row, self.materials, [])
        self.assertIsNone(res)
        self.assertEqual((conflict, blocked), ([], []))


class TestSameNameConflict(unittest.TestCase):
    """同名多物料不许自动取第一条。"""

    def test_different_codes_reported_as_conflict(self):
        materials = [mat("海信平板电视75D60S", "75D60S", "6a41aa03", code="B1"),
                     mat("海信平板电视75D60S", "75D60S", "6a41aa99", code="B2")]
        row = {"model": "75D60S", "material": "海信平板电视75D60S"}
        res, conflict, _ = mm.match_one(row, materials, [])
        self.assertIsNone(res)
        self.assertEqual(len(conflict), 2)

    def test_same_xk_code_is_one_material(self):
        materials = [mat("海信平板电视75D60S", "75D60S", "6a41aa03", code="B1"),
                     mat("海信平板电视75D60S", "75D60S", "6a41aa03", code="B2")]
        row = {"model": "75D60S", "material": "海信平板电视75D60S"}
        res, conflict, _ = mm.match_one(row, materials, [])
        self.assertIsNotNone(res, "销客编码相同应视为同一条物料，不算冲突")
        self.assertEqual(res["销客产品编码"], "6a41aa03")


class TestFuzzyGate(unittest.TestCase):
    """型号标题模糊必须过两道闸（品牌一致性 + 型号码边界）。"""

    def test_brand_mismatch_recorded_as_suspect(self):
        materials = [mat("索尼平板电视100Z600QF", "100Z600QF", "6a41aa04")]
        row = {"model": "东芝平板电视100Z600QFR-S", "material": "", "spec": ""}
        res, conflict, blocked = mm.match_one(row, materials, [])
        self.assertIsNone(res)
        self.assertTrue(blocked, "被品牌闸拦下的候选要留下来，供疑似错配清单使用")
        self.assertIn("品牌", blocked[0]["原因"])

    def test_clean_model_still_matches(self):
        materials = [mat("容声多开门冰箱BCD-606WKK1FPGZA", "BCD-606WKK1FPGZA", "6a41aa01")]
        row = {"model": "Ronshen/容声 BCD-606WKK1FPGZA", "material": "", "spec": ""}
        res, _, _ = mm.match_one(row, materials, [])
        self.assertIsNotNone(res, "干净的型号包含应能命中")
        self.assertEqual(res["匹配方式"], "型号标题模糊")

    def test_org_filter_restricts_pool(self):
        materials = [mat("容声多开门冰箱BCD-606WKK1FPGZA", "BCD-606WKK1FPGZA", "6a41aa01",
                         org="桦润芯猫"),
                     mat("容声多开门冰箱BCD-606WKK1FPGZA", "BCD-606WKK1FPGZA", "6a41aa77",
                         org="欣暖家彩电(2026)")]
        row = {"model": "BCD-606WKK1FPGZA", "material": "容声多开门冰箱BCD-606WKK1FPGZA"}
        res, _, _ = mm.match_one(row, materials, [], org="欣暖家彩电(2026)")
        self.assertEqual(res["销客产品编码"], "6a41aa77")

    def test_material_hit_without_xk_code_is_noted(self):
        """物料表命中、但该物料本身没销客编码 → 留空 + 备注，不静默。"""
        materials = [mat("某冷门物料X1", "X1", xk="")]
        row = {"model": "X1", "material": "某冷门物料X1"}
        res, _, _ = mm.match_one(row, materials, [])
        self.assertEqual(res["销客产品编码"], "")
        self.assertIn("待人工确认", res["备注"])


class TestColumnDetection(unittest.TestCase):
    """列按表头语义定位，不按列字母。"""

    @staticmethod
    def _ws(rows):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "外部清单"
        for r in rows:
            ws.append(r)
        return wb, ws

    def test_variant_headers_resolved(self):
        wb, ws = self._ws([["序号", "商品名称", "金蝶名称", "规格", "品牌"],
                           [1, "BCD-606WKK1FPGZA", "容声多开门冰箱BCD-606WKK1FPGZA",
                            "BCD-606WKK1FPGZA", "容声"]])
        cols = mm.resolve_columns(mm.read_rows(ws, 1)[0][1])
        self.assertEqual(cols["model"], 2)
        self.assertEqual(cols["material"], 3)
        self.assertEqual(cols["spec"], 4)
        self.assertEqual(cols["brand"], 5)

    def test_header_row_auto_detected_below_blank_first_row(self):
        wb, ws = self._ws([[None, None, None],
                           ["商品名称", "金蝶名称", "规格型号"],
                           ["BCD-606WKK1FPGZA", "容声多开门冰箱BCD-606WKK1FPGZA", "BCD-606WKK1FPGZA"]])
        self.assertEqual(mm.find_header_row(mm.read_rows(ws, 1, 8)), 2)

    def test_only_sheet_named_sheet1_with_real_data_is_picked(self):
        """整个文件只有一个 Sheet1 且装着真数据 → 必须用它，不能因名字跳过。"""
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["型号标题", "物料名称", "规格型号"])
        ws.append(["BCD-606WKK1FPGZA曜金印", "容声多开门冰箱BCD-606WKK1FPGZA", "BCD-606WKK1FPGZA"])
        self.assertEqual(mm.pick_sheet(wb), "Sheet1")

    def test_supply_price_sheet_without_model_column(self):
        """金蝶供货价表风格：只有 物料名称/规格型号、没有型号标题列 →
        照样能选 sheet、走第一级物料名称精确（铁律：能用物料名称就绝不用型号标题）。"""
        wb, ws = self._ws([["品牌", "物料名称", "规格型号", "供 价(RMB)", "核算价", "面价", "品类"],
                           ["美的", "美的内机MDV-D22Q4/BP", "MDV-D22Q4/BP3N1-",
                            1670, 1670, 3340, "空调"]])
        self.assertEqual(mm.pick_sheet(wb), "外部清单")
        materials = [mat("美的内机MDV-D22Q4/BP", "MDV-D22Q4/BP3N1-", "6a41aa10")]
        row = {"model": "", "material": "美的内机MDV-D22Q4/BP", "spec": "MDV-D22Q4/BP3N1-"}
        res, _, _ = mm.match_one(row, materials, [])
        self.assertEqual(res["匹配方式"], "物料名称精确")
        self.assertEqual(res["销客产品编码"], "6a41aa10")


class TestSnapshotAndMaterialTable(unittest.TestCase):
    """两个来源的「名称」不是一回事（实测踩过）：
    外部表的物料名称列与「产品主数据快照」几乎 100% 完全一致，但金蝶物料表覆盖的是另一个子集。
    所以物料名称精确要查快照 + 物料表两处；快照命中后回物料表补「销客产品编码」。"""

    def _prod(self):
        return [{"name": "容声多开门冰箱BCD-606WKK1FPGZA",
                 "mat": "容声多开门冰箱BCD-606WKK1FPGZA",
                 "spec": "BCD-606WKK1FPGZA", "code": "01.01.01.01.02.9",
                 "pclass": "", "pitem": ""}]

    def test_snapshot_hit_backfills_xk_code(self):
        materials = [mat("容声多开门冰箱BCD-606WKK1FPGZA", "BCD-606WKK1FPGZA", "6a41aa01")]
        products = self._prod()
        idx = mm.build_index(materials, products)
        row = {"model": "", "material": "容声多开门冰箱BCD-606WKK1FPGZA"}
        res, _, _ = mm.match_one(row, materials, products, index=idx)
        self.assertEqual(res["匹配方式"], "物料名称精确")
        self.assertEqual(res["来源"], "产品主数据快照")
        self.assertEqual(res["销客产品编码"], "6a41aa01", "编码要回物料表按名称/规格补回来")
        self.assertEqual(res["我方产品名称"], "容声多开门冰箱BCD-606WKK1FPGZA")

    def test_snapshot_hit_without_table_entry_leaves_code_blank(self):
        """两个来源都没有编码（旧格式快照没有唯一性ID列）→ 留空，不猜。"""
        products = self._prod()
        row = {"model": "", "material": "容声多开门冰箱BCD-606WKK1FPGZA"}
        res, _, _ = mm.match_one(row, [], products)
        self.assertEqual(res["匹配方式"], "物料名称精确")
        self.assertEqual(res["销客产品编码"], "", "物料表里没有就留空，不猜编码")

    def test_snapshot_spec_exact_is_second_tier(self):
        products = self._prod()
        row = {"model": "BCD-606WKK1FPGZA曜金印", "material": "", "spec": "BCD-606WKK1FPGZA"}
        res, _, _ = mm.match_one(row, [], products)
        self.assertEqual(res["匹配方式"], "规格型号精确")


class TestCodeFallbackFromSnapshotUid(unittest.TestCase):
    """物料表未覆盖时，销客产品编码回落用快照「唯一性ID（必填）」。

    依据（2026-09-22 用真实文件交叉验证）：快照唯一性ID 与金蝶物料表「销客产品编码」
    逐行一致 1103/1104；快照「产品编码」大面积为空、「物料编号」是 ERP 存货编码（03.05.01.02 这种），
    两者都不能当销客编码用。
    """

    UID = "6ab08cb5c1f58d0007ecc5"

    def _prod(self, uid=None):
        return [{"name": "容声多开门冰箱BCD-606WKK1FPGZA",
                 "mat": "容声多开门冰箱BCD-606WKK1FPGZA",
                 "spec": "BCD-606WKK1FPGZA", "code": "03.01.01.02.030",
                 "pclass": "冰箱", "pitem": "多门冰箱",
                 "uid": self.UID if uid is None else uid, "pcode": ""}]

    def _row(self):
        return {"model": "", "material": "容声多开门冰箱BCD-606WKK1FPGZA"}

    def test_falls_back_to_uid_when_table_missing(self):
        res, _, _ = mm.match_one(self._row(), [], self._prod())
        self.assertEqual(res["销客产品编码"], self.UID)
        self.assertIn("唯一性ID", res["备注"], "回落到快照时必须写明出处，便于抽查")

    def test_table_code_still_wins(self):
        materials = [mat("容声多开门冰箱BCD-606WKK1FPGZA", "BCD-606WKK1FPGZA", "6a41aa01")]
        res, _, _ = mm.match_one(self._row(), materials, self._prod())
        self.assertEqual(res["销客产品编码"], "6a41aa01", "物料表是官方对照，优先")
        self.assertEqual(res["备注"], "")

    def test_no_code_anywhere_stays_blank_with_note(self):
        res, _, _ = mm.match_one(self._row(), [], self._prod(uid=""))
        self.assertEqual(res["销客产品编码"], "")
        self.assertIn("待人工确认", res["备注"])

    def test_erp_material_code_not_used_as_xk_code(self):
        """物料编号（ERP 存货编码）绝不能当销客产品编码填出去。"""
        res, _, _ = mm.match_one(self._row(), [], self._prod(uid=""))
        self.assertNotIn("03.01.01.02.030", [res["销客产品编码"]])
        self.assertEqual(res["销客产品编码"], "")


class TestProductSnapshotLoader(unittest.TestCase):
    """load_products 要能取到唯一性ID列的变体表头。"""

    def test_uid_column_loaded(self):
        import build_import as bi
        tmp = tempfile.mkdtemp(prefix="prod_load_")
        try:
            p = os.path.join(tmp, "产品快照.xlsx")
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.append(["唯一性ID（必填）", "产品名称（必填）", "物料名称", "型号规格",
                       "物料编号", "产品大类", "产品品相"])
            ws.append(["6ab08cb5c1f58d0007ecc5", "容声多开门冰箱BCD-606WKK1FPGZA",
                       "容声多开门冰箱BCD-606WKK1FPGZA", "BCD-606WKK1FPGZA",
                       "03.01.01.02.030", "冰箱", "多门冰箱"])
            wb.save(p)
            prods = bi.load_products(p)
            self.assertEqual(len(prods), 1)
            self.assertEqual(prods[0]["uid"], "6ab08cb5c1f58d0007ecc5")
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


class TestEndToEnd(unittest.TestCase):
    """端到端：生成三 sheet 报告，未匹配原因分类正确。"""

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="match_test_")
        self.ext = os.path.join(self.tmp, "外部清单.xlsx")
        self.mt = os.path.join(self.tmp, "物料_9999999999999_190792.xlsx")
        self.snap = os.path.join(self.tmp, "产品快照.xlsx")
        self.outdir = os.path.join(self.tmp, "out")
        self._make_ext()
        self._make_mt()
        self._make_snapshot()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _make_ext(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "货品清单"
        ws.append(["品牌", "型号标题", "物料名称", "规格型号"])
        ws.append(["Ronshen/容声", "BCD-606WKK1FPGZA曜金印",
                   "容声多开门冰箱BCD-606WKK1FPGZA", "BCD-606WKK1FPGZA"])
        ws.append(["TCL", "TCL平板电视98C10L-A", "0", "98C10L-A"])
        ws.append(["海尔", "某不知名热水器X999", "", ""])
        # 只在快照里、物料表没覆盖的一条：用来验证唯一性ID 回填能一路走到报告里
        ws.append(["海信", "海信激光电视100L6N流砂锖", "海信激光电视100L6N流砂锖", "100L6N"])
        wb.create_sheet("Sheet2").append(["草稿"])
        wb.save(self.ext)

    def _make_mt(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["使用组织", "编码", "名称", "规格型号", "销客产品编码"])
        ws.append(["桦润芯猫", "01.01.01.01.02.001", "容声多开门冰箱BCD-606WKK1FPGZA",
                   "BCD-606WKK1FPGZA", "6a41aa01"])
        ws.append(["桦润芯猫", "01.01.01.01.02.002", "TCL平板电视98C10L-A",
                   "98C10L-A", "6a41aa02"])
        wb.save(self.mt)

    def _make_snapshot(self):
        """给测试一个隔离的产品快照，避免读到机器上真实快照。
        带唯一性ID 列，与真实产品导出结果一致。"""
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["唯一性ID（必填）", "产品名称（必填）", "物料名称", "型号规格",
                   "物料编号", "产品大类", "产品品相"])
        ws.append(["6a00000000000000000001", "与本次无关的产品X", "与本次无关的产品X",
                   "XXXX", "0", "", ""])
        ws.append(["6ab08cb5c1f58d0007ecc5", "海信激光电视100L6N流砂锖",
                   "海信激光电视100L6N流砂锖", "100L6N", "03.05.01.01", "彩电", "激光电视"])
        wb.save(self.snap)

    def _run(self):
        cmd = [sys.executable, os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                            "match_material.py"),
               "--source", self.ext, "--material-table", self.mt,
               "--product-snapshot", self.snap, "--downloads", self.tmp,
               "--outdir", self.outdir]
        r = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8")
        return r

    def test_report_sheets_and_counts(self):
        r = self._run()
        self.assertEqual(r.returncode, 0, r.stderr)
        payload = json.loads(r.stdout[r.stdout.index("{"):])
        self.assertEqual(payload["sheet"], "货品清单", "要跳过 Sheet2，选中真数据表")
        self.assertEqual(payload["数据行"], 4)
        self.assertEqual(payload["匹配上"], 3)
        self.assertEqual(payload["未匹配"], 1)
        self.assertEqual(payload["匹配方式分布"].get("物料名称精确"), 2)
        self.assertEqual(payload["匹配方式分布"].get("规格型号精确"), 1)
        self.assertEqual(payload["未匹配原因分布"].get("我方无此物料"), 1)
        self.assertEqual(payload["已匹配但取不到销客产品编码"], 0,
                         "物料表 + 快照唯一性ID 两路下来不该再有取不到编码的行")

        from openpyxl import load_workbook
        wb = load_workbook(payload["files"][0], data_only=True)
        self.assertEqual(wb.sheetnames, ["匹配结果", "未匹配清单", "统计"])
        ws = wb["匹配结果"]
        hdr = [c.value for c in ws[1]]
        for col in ["我方产品名称", "我方物料名称", "物料编码", "规格型号",
                    "销客产品编码", "所属组织", "匹配方式", "置信度"]:
            self.assertIn(col, hdr)
        self.assertEqual(ws.max_row, 5)
        # 只在快照里、物料表没覆盖的那一行：编码应来自快照唯一性ID
        i_mat, i_xk = hdr.index("我方物料名称") + 1, hdr.index("销客产品编码") + 1
        i_note, i_src = hdr.index("备注") + 1, hdr.index("来源") + 1
        hit = [r for r in range(2, ws.max_row + 1)
               if ws.cell(r, i_mat).value == "海信激光电视100L6N流砂锖"]
        self.assertEqual(len(hit), 1, "快照命中的那一行必须出现在报告里")
        r = hit[0]
        self.assertEqual(ws.cell(r, i_xk).value, "6ab08cb5c1f58d0007ecc5")
        self.assertEqual(ws.cell(r, i_src).value, "产品主数据快照")
        self.assertIn("唯一性ID", ws.cell(r, i_note).value or "")
        ws2 = wb["未匹配清单"]
        self.assertIn("我方无此物料", [ws2.cell(r, 1).value for r in range(2, ws2.max_row + 1)])

    def test_missing_material_table_stops(self):
        """两个底表都没有 → 停手，不做无依据的猜测匹配。"""
        cmd = [sys.executable, os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                            "match_material.py"),
               "--source", self.ext, "--material-table", os.path.join(self.tmp, "没有这个.xlsx"),
               "--downloads", self.tmp, "--outdir", self.outdir]
        r = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8")
        self.assertNotEqual(r.returncode, 0, "一个底表都没有时必须停手，不能猜")

    def test_missing_material_table_degrades_to_snapshot(self):
        """有产品快照、缺物料表 → 降级继续跑，编码取自快照唯一性ID，并在结果里写明降级。"""
        cmd = [sys.executable, os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                            "match_material.py"),
               "--source", self.ext, "--material-table", os.path.join(self.tmp, "没有这个.xlsx"),
               "--product-snapshot", self.snap, "--downloads", self.tmp, "--outdir", self.outdir]
        r = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(r.returncode, 0, r.stderr)
        payload = json.loads(r.stdout[r.stdout.index("{"):])
        self.assertIn("降级说明", payload, "降级必须写进结果，不能让用户以为是正常跑")
        self.assertIn("物料表", payload["降级说明"])
        self.assertEqual(payload["匹配上"], 1, "只有快照里那条能匹配上")
        self.assertEqual(payload["已匹配但取不到销客产品编码"], 0)

        from openpyxl import load_workbook
        wb = load_workbook(payload["files"][0], data_only=True)
        ws = wb["匹配结果"]
        hdr = [c.value for c in ws[1]]
        i_xk = hdr.index("销客产品编码") + 1
        codes = [ws.cell(r0, i_xk).value for r0 in range(2, ws.max_row + 1)]
        self.assertIn("6ab08cb5c1f58d0007ecc5", codes)


if __name__ == "__main__":
    unittest.main(verbosity=2)
