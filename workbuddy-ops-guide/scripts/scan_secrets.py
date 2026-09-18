# -*- coding: utf-8 -*-
"""
敏感词扫描（对外文件交付前必跑）

★ 三条要点
1. 打印命中行的【上下文原文】，由人工判断是否误报（命中常在免责句里）
2. Office 文件是压缩包，必须取真文本后再扫——直接扫二进制必然假阴性
3. 取不到文本时明确报警，绝不静默报"零命中"

用法:
    python scan_secrets.py <文件或目录> [额外关键词,逗号分隔]

返回码:
    0 零命中   1 有命中   2 用法错误或路径不存在   3 有文件未取到文本

★ 默认词表是跨行业通用的内部口径词，不含任何具体主体信息。
  把你自己的平台名、供应商名、内部对象前缀用第二个参数追加：
      python scan_secrets.py a.docx "我的平台名,供应商A,内部对象前缀"
"""
import os
import sys
import zipfile
from html import unescape

DEFAULT_WORDS = [
    # 费用与议价（内部口径）
    '人天', '核减', '自做', '议价', '压价', '砍价', '报价单',
    # 我方内部判定
    '可自做', '建议核减', '议价点',
    # 选型（对外一般不提）
    '替换', '换平台', '更换系统',
]

TEXT_EXT = {'.md', '.txt', '.py', '.js', '.ts', '.json', '.csv', '.groovy',
            '.yaml', '.yml', '.xml', '.html', '.htm', '.sh', '.sql',
            '.ini', '.cfg', '.log', '.rst'}
OFFICE_EXT = {'.docx', '.xlsx', '.xlsm', '.pptx'}
SKIP_DIRS = {'.git', '__pycache__', '_backups', 'node_modules', '.venv', '.idea'}


def _zip_xml_text(path, prefixes=None):
    """兜底：解压 OOXML 取 XML 文本。
    ★ 中文常被写成数字实体（如 &#20154;），必须 unescape 才搜得到。"""
    try:
        z = zipfile.ZipFile(path)
    except Exception:
        return None
    buf = []
    try:
        for name in z.namelist():
            if not name.endswith('.xml'):
                continue
            if prefixes and not name.startswith(prefixes):
                continue
            buf.append(unescape(z.read(name).decode('utf-8', 'ignore')))
    except Exception:
        return None
    return '\n'.join(buf) if buf else None


def _docx_text(path):
    try:
        import docx
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                for c in row.cells:
                    parts.append(c.text)
        return '\n'.join(parts)
    except Exception:
        return _zip_xml_text(path, ('word/',))


def _xlsx_text(path):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
        parts = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for v in row:
                    if v is not None:
                        parts.append(str(v))
        return '\n'.join(parts)
    except Exception:
        return _zip_xml_text(path, ('xl/',))


def _pptx_text(path):
    return _zip_xml_text(path, ('ppt/',))


def extract_text(path):
    """返回 (文本, 状态)；状态: ok / no-text / error"""
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext in TEXT_EXT:
            for enc in ('utf-8', 'gbk', 'utf-16'):
                try:
                    return open(path, encoding=enc).read(), 'ok'
                except (UnicodeDecodeError, UnicodeError):
                    continue
            return open(path, 'rb').read().decode('utf-8', 'ignore'), 'ok'
        if ext == '.docx':
            t = _docx_text(path)
        elif ext in ('.xlsx', '.xlsm'):
            t = _xlsx_text(path)
        elif ext == '.pptx':
            t = _pptx_text(path)
        else:
            return '', 'no-text'
    except Exception:
        return '', 'error'
    if t is None:
        return '', 'no-text'
    return t, 'ok'


def scan_file(path, words):
    """返回 (命中列表, 状态)"""
    text, status = extract_text(path)
    if status != 'ok':
        return [], status
    hits = []
    for i, line in enumerate(text.split('\n'), 1):
        for w in words:
            if w in line:
                hits.append((i, w, line.strip()[:160]))
                break
    return hits, 'ok'


def collect(path):
    if os.path.isfile(path):
        return [path]
    out = []
    for root, dirs, files in os.walk(path):
        dirs[:] = [d for d in dirs if d not in SKIP_DIRS]
        for f in files:
            if f.startswith('~$'):
                continue
            out.append(os.path.join(root, f))
    return sorted(out)


def main(argv):
    if len(argv) < 2:
        print(__doc__)
        return 2
    target = argv[1]
    extra = argv[2] if len(argv) > 2 else None

    words = list(DEFAULT_WORDS)
    if extra:
        words += [w.strip() for w in extra.split(',') if w.strip()]

    if not os.path.exists(target):
        print('路径不存在:', target)
        return 2

    files = collect(target)
    if not files:
        print('目录下无可扫描文件:', target)
        return 2

    total_hits = 0
    no_text = []
    err_files = []

    print('=' * 70)
    print('扫描目标:', target)
    print('文件数:', len(files), ' 扫描词数:', len(words))
    print('=' * 70)

    for p in files:
        hits, status = scan_file(p, words)
        if status == 'no-text':
            no_text.append(p)
            print('  [跳过] 未取到文本:', os.path.basename(p))
            continue
        if status == 'error':
            err_files.append(p)
            print('  [错误] 读取失败:', os.path.basename(p))
            continue
        if hits:
            total_hits += len(hits)
            print('\n  命中 %s  %d 处:' % (os.path.basename(p), len(hits)))
            for ln, w, ctx in hits:
                print('      [行%d] 词「%s」' % (ln, w))
                print('             %s' % ctx)
        else:
            print('  [干净] %s' % os.path.basename(p))

    print()
    print('-' * 70)
    if total_hits:
        print('合计命中 %d 处（必须逐条看上下文判断是否为误报）' % total_hits)
        print('  - 出现在「本件不含…」等免责/说明句里 → 误报，保留')
        print('  - 出现在正文描述里 → 真泄漏，必须改')
    else:
        print('零命中')
    if no_text:
        print('以下文件未取到文本，不能视为已检查，需人工确认：')
        for p in no_text:
            print('     ', os.path.basename(p))
    if err_files:
        print('读取失败（需人工确认）：')
        for p in err_files:
            print('     ', os.path.basename(p))
    print('-' * 70)

    if total_hits:
        return 1
    if no_text or err_files:
        return 3
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv))
