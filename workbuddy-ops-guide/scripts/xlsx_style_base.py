# -*- coding: utf-8 -*-
"""
xlsx 生成样式基线 + 安全保存（含占用另存）
本机统一风格：微软雅黑 / 表头 4472C4 / 冻结首行 / 自动换行

用法: 从本文件 import 后使用，或复制 sheet() 函数改造
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = '微软雅黑'

H_FILL = PatternFill('solid', fgColor='4472C4')      # 表头 蓝
SUB_FILL = PatternFill('solid', fgColor='2E5C8A')    # 次级表头 深蓝
SEC_FILL = PatternFill('solid', fgColor='D9E2F3')    # 分组行 浅蓝
TIP_FILL = PatternFill('solid', fgColor='FFF2CC')    # 提示 黄
WARN_FILL = PatternFill('solid', fgColor='FCE4D6')   # 待确认 橙
OK_FILL = PatternFill('solid', fgColor='E2EFDA')     # 已完成 绿
GREY_FILL = PatternFill('solid', fgColor='F2F2F2')   # 待填 灰

THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def safe_save(wb, path):
    """★ 写盘前探测占用，占用则 _v2/_v3 递增另存，返回实际路径"""
    try:
        open(path, 'a').close()
    except PermissionError:
        base, ext = os.path.splitext(path)
        i = 2
        while True:
            alt = f'{base}_v{i}{ext}'
            try:
                open(alt, 'a').close()
                path = alt
                break
            except PermissionError:
                i += 1
    wb.save(path)
    return path


def sheet(wb, name, headers, rows, widths, note=None, group_col=None,
          state_col=None, freeze=True):
    """
    headers: 表头列表
    rows:    数据行；以 '§' 开头的第一格 = 分组小标题行
    widths:  列宽列表
    group_col: 分组行合并到第几列
    state_col: 该列(1-based)按内容自动着色（待确认橙 / 已完成绿 / 空灰）
    """
    ws = wb.create_sheet(name)
    r = 1
    if note:
        ws.cell(row=1, column=1, value=note).font = Font(name=FONT, size=9,
                                                         italic=True, color='7F7F7F')
        ws.merge_cells(start_row=1, start_column=1, end_row=1,
                       end_column=max(len(headers), 2))
        ws.row_dimensions[1].height = 30
        ws.cell(row=1, column=1).alignment = Alignment(wrap_text=True, vertical='center')
        r = 2

    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
        cell.fill = H_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[r].height = 28
    head_row = r
    r += 1

    for row in rows:
        if row and isinstance(row[0], str) and row[0].startswith('§'):
            ws.cell(row=r, column=1, value=row[0][1:]).font = Font(
                name=FONT, size=10, bold=True, color='1F3864')
            for c in range(1, len(headers) + 1):
                ws.cell(row=r, column=c).fill = SEC_FILL
                ws.cell(row=r, column=c).border = BORDER
            if group_col:
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=group_col)
            r += 1
            continue
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            cell.border = BORDER
        if state_col and state_col <= len(row):
            sv = str(row[state_col - 1] or '')
            if sv.startswith(('待确认', '待核对')) or '待补' in sv or '缺口' in sv:
                ws.cell(row=r, column=state_col).fill = WARN_FILL
            elif sv.startswith(('已确认', '在用', '已实现', '已上线')):
                ws.cell(row=r, column=state_col).fill = OK_FILL
            elif not sv or sv.startswith('（待填'):
                ws.cell(row=r, column=state_col).fill = GREY_FILL
        r += 1

    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    if freeze:
        ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    return ws


def unify_number_format(ws, col_letter, fmt, start_row=2):
    """★ 生成后按列强制统一 number_format，防百分比/小数混排"""
    for row in ws[col_letter]:
        if row.row >= start_row:
            row.number_format = fmt


if __name__ == '__main__':
    # 示例
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheet(wb, '示例',
          ['#', '事项', '状态', '核对结果'],
          [['§一、进行中'],
           ['1', '某某单据主对象 API', '待确认', ''],
           ['2', '跟进记录对象 API', '已确认', 'ExampleRecordObj']],
          [6, 30, 12, 26],
          note='示例：分组行用 § 开头，状态列自动着色',
          group_col=4, state_col=3)
    p = safe_save(wb, os.path.join(os.path.dirname(__file__), '_demo.xlsx'))
    print('SAVED:', p)
