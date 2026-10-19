# -*- coding: utf-8 -*-
"""
欣暖家月度考核金额表生成器 v2.0（通用版）
================================================
换月只需丢入「本月基础表 + 上期成品（当模板）+ 最新人员导出」，其余自动推导。

自动推导：
  1. 部门块顺序      —— 按基础表「二级部门」首次出现顺序，父块(美的特渠部)展开为子块
  2. 部门别名        —— 天猫彩电店 -> 天猫彩电部
  3. 跳过部门        —— 天猫零售部（用户指定某月不做）
  4. 人员目标        —— 目标表按月份过滤（试用/正式自动切换）；查不到则回落到报表「指标」列
  5. 小计目标        —— 成员目标相加（8 月定案口径）；也可切回目标表合计行
  6. 小计行标签      —— 门店=「XX合计」、渠道=「渠道N部小计」、父块=「渠道部合计」
  7. 样式            —— 自动扫描模板金额表里的「人员行 / 小计行」样式与数字格式
  8. 人员表          —— 可合并多份人员导出（按唯一性ID去重）
  9. 目标表          —— 按月重建（保留生效区间标签）+ 合计行重算

用法：
  python gen_xinnuanjia_amount.py --base 基础表.xlsx --template 上期成品.xlsx \
         --month 9 --out 欣暖家9月考核金额表.xlsx \
         --roster 人员导出1.xlsx [--roster 人员导出2.xlsx ...]

可选：
  --config cfg.json          覆盖默认配置（部门别名/跳过/人员归部/父子/顺序）
  --target-source sheet|report     目标来源，默认 sheet（目标表优先，缺失回落报表）
  --subtotal-from members|sheet    小计目标来源，默认 members（成员相加）
  --keep-target              不重建目标表（原样保留模板的）
  --skip-dept NAME ...       追加跳过部门；--skip-person NAME ... 追加跳过人员
  --person-dept "姓名:部门"   追加人员归部覆盖（特渠分部等）
  --allow-unknown-dept       允许出现未登记部门（默认报错，禁止猜）
"""
import argparse, json, re
from copy import copy
import openpyxl

# ------------------------------------------------------------------ 默认配置
DEFAULT_CONFIG = {
    # 基础表「二级部门」-> 输出部门名
    "alias": {"天猫彩电店": "天猫彩电部"},
    # 已登记部门（白名单，按 alias 后的名字）。基础表出现名单外的部门 -> 报错，禁止静默纳入
    "known_depts": ["美的滨江店", "东芝华东店", "天猫彩电部", "东芝宝龙店",
                    "特定渠道1部", "特定渠道2部", "美的特渠部", "天猫零售部"],
    # 本月不做的部门（整块跳过）
    "skip_depts": [],
    # 整块跳过的人员（部门仍在，只剔人）
    "skip_people": [],
    # 人员 -> 输出部门（用于把父部门拆成 1部/2部）
    "person_dept": {
        "吴建": "特定渠道1部", "车邦建": "特定渠道1部",
        "韩芬": "特定渠道1部", "周晨烨": "特定渠道1部",
        "王浩萍": "特定渠道2部", "郁轹文": "特定渠道2部",
        "钟俊楠": "特定渠道2部", "陈学健": "特定渠道2部",
    },
    # 父块：块名 -> 子块列表（父块自身不出人员，只出汇总小计）
    "parent_blocks": {"美的特渠部": ["特定渠道1部", "特定渠道2部"]},
    # 小计行显示名（未列出的用「XX合计」）
    "sub_label": {"特定渠道1部": "渠道1部小计", "特定渠道2部": "渠道2部小计",
                  "美的特渠部": "渠道部合计"},
    # 块内人员顺序（未列出的块按基础表原序；不在名单里的排最后）
    "name_order": {
        "特定渠道1部": ["吴建", "车邦建", "韩芬", "周晨烨"],
        "特定渠道2部": ["王浩萍", "郁轹文", "钟俊楠", "陈学健"],
    },
    # 奖金基数
    "bonus_person": 300,
    "bonus_sub_monthly": 2000,
    "bonus_sub_other": 1000,
}

OUT_HEADERS = ["主属部门-3", "员工姓名", "达成奖金",
               "月度销售指标", "月度整体销售金额（含变更）", "月度整体达成率", "月度整体达成奖金",
               "第一成交率金额指标", "第一成交率金额（含变更）", "第一成交率金额达成率", "第一成交达成奖金",
               "ABC成交金额指标", "ABC成交金额（含变更）", "ABC成交金额达成率", "ABC成交达成奖金",
               "ABC成交客户数（订单对象统计）", "ABC客单价",
               "ABC成交顾客数指标", "ABC成交顾客数", "ABC成交顾客数达成率", "ABC成交顾客数达成奖金",
               "录入ABC线索数指标", "录入ABC线索数", "ABC线索新建达成率", "ABC录入线索达成奖金"]

AMOUNT_COLS = [3, 4, 5, 7, 8, 9, 11, 12, 13, 15, 17, 21, 25]
INT_COLS = [16, 18, 19, 22, 23]


def num(v):
    return v if isinstance(v, (int, float)) else 0.0


def bare(nm):
    return re.sub(r"[（(].*?[)）]", "", str(nm)).strip()


def parse_period(txt):
    """（7月试用期）-> {7}；（8-12月）-> {8..12}；（7.8月试用期）-> {7,8}；其余 -> 全年"""
    if not txt:
        return set(range(1, 13))
    t = str(txt).strip("（）()")
    m = re.search(r"(\d+)(?:\.(\d+))?月试用期", t)
    if m:
        return {int(m.group(1)), int(m.group(2))} if m.group(2) else {int(m.group(1))}
    m = re.search(r"(\d+)-(\d+)月", t)
    if m:
        return set(range(int(m.group(1)), int(m.group(2)) + 1))
    return set(range(1, 13))


# ------------------------------------------------------------------ 读入
def load_base(path, cfg):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise SystemExit(f"[X] 基础表为空: {path}")
    persons, dept_seq = [], []
    cur = ""          # 二级部门只在块首行出现，后续行为空 -> 显式向下填充
    for r in rows[1:]:
        if r[2] is None:
            continue
        ns = str(r[2]).strip()
        if ns == "" or "小计" in ns or "合计" in ns or ns == "总计":
            continue
        d = str(r[1] or "").strip()
        if d:
            cur = d
        else:
            d = cur
        if d and d not in dept_seq:
            dept_seq.append(d)
        if bare(ns) in cfg["skip_people"]:
            continue
        persons.append({
            "name": ns, "base_dept": d,
            "monthly": num(r[5]), "first_deal": num(r[9]), "abc": num(r[13]),
            "abc_cust": num(r[15]), "abc_cust_n": num(r[18]), "leads": num(r[21]),
            "r_monthly": num(r[3]), "r_first": num(r[7]), "r_abc": num(r[11]),
            "r_cust": num(r[17]), "r_leads": num(r[20]),
            "disabled": "停用" in ns,
        })
    return persons, dept_seq


def load_target_sheet(path, month):
    wb = openpyxl.load_workbook(path, data_only=True)
    if "目标表" not in wb.sheetnames:
        wb.close()
        return {}, {}
    ws = wb["目标表"]
    per, sums = {}, {}
    for r in range(1, ws.max_row + 1):
        who = ws.cell(r, 2).value
        store = str(ws.cell(r, 1).value or "").strip()
        if not who:
            continue
        who = str(who).strip()
        is_sub = ("合计" in store or "小计" in store
                  or "合计" in who or "小计" in who or who == "人员")
        vals = {"monthly": num(ws.cell(r, 4).value), "first_deal": num(ws.cell(r, 5).value),
                "abc": num(ws.cell(r, 7).value), "cust": num(ws.cell(r, 9).value),
                "leads": num(ws.cell(r, 11).value)}
        if is_sub:
            if store and store != "人员":
                sums[store] = vals
            continue
        if month not in parse_period(ws.cell(r, 3).value):
            continue
        per[who] = dict(vals, period=ws.cell(r, 3).value)
    wb.close()
    return per, sums


# ------------------------------------------------------------------ 部门块
def derive_blocks(dept_seq, cfg):
    alias, skip = cfg["alias"], set(cfg["skip_depts"])
    parents = cfg["parent_blocks"]
    blocks, order = [], []
    for d in dept_seq:
        d2 = alias.get(d, d)
        if d2 in skip:
            continue
        for b in parents.get(d2, [d2]):
            if b not in blocks:
                blocks.append(b)
            if b not in order:
                order.append(b)
    # 父块放到最后（只出汇总小计）
    return [b for b in order if b not in parents] + [b for b in order if b in parents]


def assign(persons, blocks, cfg):
    """归部。返回 (groups, unknown, in_parent, skipped)

    skip_depts 里的人整块丢弃（既不进 groups，也不算 unknown）。
    """
    alias, parents, pd = cfg["alias"], cfg["parent_blocks"], cfg["person_dept"]
    skip = set(cfg["skip_depts"])
    groups, unknown, in_parent, skipped = {}, [], [], []
    for p in persons:
        b = bare(p["name"])
        blk = pd[b] if b in pd else alias.get(p["base_dept"], p["base_dept"])
        if blk in skip:                      # 整块跳过的部门 -> 该人员不参与本月核算
            skipped.append((p["name"], blk))
            continue
        if blk in parents:
            in_parent.append(p["name"])
            blk = "__未分配__"
        if blk not in blocks and blk != "__未分配__":
            unknown.append((p["name"], blk))
        p["dept"] = blk
        groups.setdefault(blk, []).append(p)
    return groups, unknown, in_parent, skipped


# ------------------------------------------------------------------ 目标与奖金
def resolve_targets(persons, sheet_tgt, source):
    missing, fallback = [], []
    for p in persons:
        b = bare(p["name"])
        if source == "sheet" and b in sheet_tgt:
            t = sheet_tgt[b]
            p["T"] = {k: t[k] for k in ("monthly", "first_deal", "abc", "cust", "leads")}
            p["period"], p["tsrc"] = t["period"], "目标表"
        elif p["disabled"]:
            p["T"] = dict(monthly=0, first_deal=0, abc=0, cust=0, leads=0)
            p["period"], p["tsrc"] = None, "停用"
        else:
            p["T"] = {"monthly": p["r_monthly"], "first_deal": p["r_first"], "abc": p["r_abc"],
                      "cust": p["r_cust"], "leads": p["r_leads"]}
            p["period"], p["tsrc"] = None, "报表指标"
            fallback.append(p["name"])
            if p["T"]["monthly"] == 0:
                missing.append(p["name"])
    return missing, fallback


def bonus_person(p, cfg):
    if p["disabled"]:
        return [0, 0, 0, 0, 0, 0]
    B, t = cfg["bonus_person"], p["T"]
    rt = lambda a, b: (a / b) if b else 0.0
    b4 = B if rt(p["monthly"], t["monthly"]) >= 1 else 0
    b1 = B if rt(p["first_deal"], t["first_deal"]) >= 1 else 0
    b3a = B if rt(p["abc"], t["abc"]) >= 1 else 0
    b3c = B if rt(p["abc_cust_n"], t["cust"]) >= 1 else 0
    b2 = B if rt(p["leads"], t["leads"]) >= 1 else 0
    return [b4, b1, b3a, b3c, b2, b4 + b1 + max(b3a, b3c) + b2]


def sum_recs(recs):
    s = dict(monthly=0.0, first_deal=0.0, abc=0.0, abc_cust=0.0, abc_cust_n=0.0,
             leads=0.0, n=len(recs))
    for x in recs:
        for k in ("monthly", "first_deal", "abc", "abc_cust", "abc_cust_n", "leads"):
            s[k] += x[k]
    for k in ("monthly", "first_deal", "abc", "cust", "leads"):
        s["T_" + k] = sum(x["T"][k] for x in recs)
    return s


# ------------------------------------------------------------------ 行
def prow(p, cfg):
    t = p["T"]
    b4, b1, b3a, b3c, b2, tot = bonus_person(p, cfg)
    cu = p["abc_cust"]
    return [p["dept"], p["name"], tot, t["monthly"], p["monthly"], None, b4,
            t["first_deal"], p["first_deal"], None, b1,
            t["abc"], p["abc"], None, b3a,
            p["abc_cust"], (p["abc"] / cu if cu else 0.0),
            t["cust"], p["abc_cust_n"], None, b3c,
            t["leads"], p["leads"], None, b2]


def srow(dept, recs, cfg, sheet_sum=None):
    s = sum_recs(recs)
    if sheet_sum:
        T = {k: sheet_sum[k] for k in ("monthly", "first_deal", "abc", "cust", "leads")}
    else:
        T = {k: s["T_" + k] for k in ("monthly", "first_deal", "abc", "cust", "leads")}
    rt = lambda a, b: (a / b) if b else 0.0
    b4 = cfg["bonus_sub_monthly"] if rt(s["monthly"], T["monthly"]) >= 1 else 0
    b1 = cfg["bonus_sub_other"] if rt(s["first_deal"], T["first_deal"]) >= 1 else 0
    b3a = cfg["bonus_sub_other"] if rt(s["abc"], T["abc"]) >= 1 else 0
    b3c = cfg["bonus_sub_other"] if rt(s["abc_cust_n"], T["cust"]) >= 1 else 0
    b2 = cfg["bonus_sub_other"] if rt(s["leads"], T["leads"]) >= 1 else 0
    tot = b4 + b1 + b2 + (cfg["bonus_sub_other"] if (b3a + b3c) >= cfg["bonus_sub_other"] else 0)
    cu = s["abc_cust"]
    return [dept, "小计", tot, T["monthly"], s["monthly"], None, b4,
            T["first_deal"], s["first_deal"], None, b1,
            T["abc"], s["abc"], None, b3a,
            s["abc_cust"], (s["abc"] / cu if cu else 0.0),
            T["cust"], s["abc_cust_n"], None, b3c,
            T["leads"], s["leads"], None, b2]


def ordered(recs, dept, cfg):
    ordr = cfg["name_order"].get(dept)
    if not ordr:
        return recs
    o = {n: i for i, n in enumerate(ordr)}
    return sorted(recs, key=lambda p: o.get(bare(p["name"]), 999))


def make_rows(blocks, groups, cfg, sheet_sums, subtotal_from):
    rows = []
    for dept in blocks:
        recs = ordered(groups.get(dept, []), dept, cfg)
        for p in recs:
            rows.append(prow(p, cfg))
        key = cfg["sub_label"].get(dept, dept + "合计")
        ss = sheet_sums.get(key) if subtotal_from == "sheet" else None
        rows.append(srow(dept, recs, cfg, ss))
    for parent, kids in cfg["parent_blocks"].items():
        recs = []
        for k in kids:
            recs += groups.get(k, [])
        if not recs:
            continue
        key = cfg["sub_label"].get(parent, parent + "合计")
        ss = sheet_sums.get(key) if subtotal_from == "sheet" else None
        rows.append(srow(parent, recs, cfg, ss))
    return rows


# ------------------------------------------------------------------ 写工作簿
def cs(src, dst):
    dst.font = copy(src.font); dst.border = copy(src.border); dst.fill = copy(src.fill)
    dst.alignment = copy(src.alignment); dst.number_format = src.number_format
    dst.protection = copy(src.protection)


def find_style_rows(ws):
    person = sub = None
    for r in range(2, ws.max_row + 1):
        b = ws.cell(r, 2).value
        if b is None:
            continue
        b = str(b).strip()
        if b in ("小计", "合计"):
            sub = sub or r
        elif "小计" not in b and "合计" not in b:
            person = person or r
    return person or 2, sub or 2


def build_amount_sheet(wb, tmpl, rows, title):
    src = openpyxl.load_workbook(tmpl)[title]
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title, 0)
    ws.append(OUT_HEADERS)
    for c in range(1, 26):
        cs(src.cell(1, c), ws.cell(1, c))
    p_rep, s_rep = find_style_rows(src)
    r = 2
    for row in rows:
        for c in range(1, 26):
            ws.cell(r, c, row[c - 1])
        rep = p_rep if row[1] != "小计" else s_rep
        for c in range(1, 26):
            cs(src.cell(rep, c), ws.cell(r, c))
        r += 1
    last = r - 1
    for row in range(2, last + 1):
        sub = (ws.cell(row, 2).value == "小计")
        ws.cell(row, 6, f"=IF(D{row}=0,0,E{row}/D{row})")
        ws.cell(row, 10, f"=IF(H{row}=0,0,I{row}/H{row})")
        ws.cell(row, 14, f"=IF(L{row}=0,0,M{row}/L{row})")
        ws.cell(row, 20, f"=IF(R{row}=0,0,S{row}/R{row})")
        ws.cell(row, 24, f"=IF(V{row}=0,0,W{row}/V{row})")
        base, thr = (2000, 1000) if sub else (300, 300)
        ws.cell(row, 7, f"=IF(F{row}>=1,{base},0)")
        ws.cell(row, 11, f"=IF(J{row}>=1,{thr},0)")
        ws.cell(row, 15, f"=IF(N{row}>=1,{thr},0)")
        ws.cell(row, 21, f"=IF(T{row}>=1,{thr},0)")
        ws.cell(row, 25, f"=IF(X{row}>=1,{thr},0)")
        ws.cell(row, 3, f"=G{row}+K{row}+Y{row}+IF((O{row}+U{row})>={thr},{thr},0)")
    for row in range(2, last + 1):
        for c in AMOUNT_COLS:
            ws.cell(row, c).number_format = "#,##0.00"
        for c in INT_COLS:
            ws.cell(row, c).number_format = "0"
    for col, dim in src.column_dimensions.items():
        if dim.width:
            ws.column_dimensions[col].width = dim.width
    ws.freeze_panes = src.freeze_panes
    src.parent.close()
    return ws, last


def rebuild_target_sheet(wb, tmpl, blocks, groups, cfg):
    ref = openpyxl.load_workbook(tmpl)["目标表"]
    ws = wb["目标表"]
    hdr = None
    for r in range(1, min(ref.max_row, 40) + 1):
        if str(ref.cell(r, 2).value or "").strip() == "人员":
            hdr = r
            break
    hdr = hdr or 3
    p_rep = s_rep = None
    for r in range(hdr + 1, ref.max_row + 1):
        a = str(ref.cell(r, 1).value or "")
        if "合计" in a or "小计" in a:
            s_rep = s_rep or r
        elif ref.cell(r, 2).value:
            p_rep = p_rep or r
    p_rep = p_rep or hdr + 1
    s_rep = s_rep or p_rep
    sty_p = [copy(ref.cell(p_rep, c)._style) for c in range(1, 13)]
    sty_s = [copy(ref.cell(s_rep, c)._style) for c in range(1, 13)]
    for r in range(hdr + 1, max(ws.max_row, hdr + 1) + 60):
        for c in range(1, 15):
            ws.cell(r, c).value = None
    r = hdr + 1

    def put_sub(label, store, s, n_people):
        nonlocal r
        vals = [label, store, None, s["T_monthly"], s["T_first_deal"], n_people * cfg["bonus_person"],
                s["T_abc"], None, s["T_cust"], n_people * cfg["bonus_person"],
                s["T_leads"], n_people * cfg["bonus_person"]]
        for c in range(1, 13):
            ws.cell(r, c, vals[c - 1])._style = copy(sty_s[c - 1])
        r += 1

    for dept in blocks:
        recs = ordered(groups.get(dept, []), dept, cfg)
        for p in recs:
            t = p["T"]
            vals = [dept, bare(p["name"]), p["period"], t["monthly"], t["first_deal"],
                    cfg["bonus_person"], t["abc"], None, t["cust"], cfg["bonus_person"],
                    t["leads"], cfg["bonus_person"]]
            for c in range(1, 13):
                ws.cell(r, c, vals[c - 1])._style = copy(sty_p[c - 1])
            r += 1
        put_sub(cfg["sub_label"].get(dept, dept + "合计"), dept, sum_recs(recs), len(recs))
    for parent, kids in cfg["parent_blocks"].items():
        recs = []
        for k in kids:
            recs += groups.get(k, [])
        if not recs:
            continue
        put_sub(cfg["sub_label"].get(parent, parent + "合计"), "", sum_recs(recs), len(recs))
    return r - 1


def replace_personnel(wb, paths):
    srcs = []
    for p in paths:
        w = openpyxl.load_workbook(p, data_only=True)
        nm = "人员数据" if "人员数据" in w.sheetnames else w.sheetnames[0]
        srcs.append(w[nm])
    idx = wb.sheetnames.index("人员表") if "人员表" in wb.sheetnames else 1
    if "人员表" in wb.sheetnames:
        del wb["人员表"]
    ws = wb.create_sheet("人员表", idx)
    seen, rr = set(), 0
    for i, s in enumerate(srcs):
        for row in s.iter_rows():
            key = row[0].value
            if i > 0 and key in seen:
                continue
            seen.add(key)
            rr += 1
            for c in row:
                if c.value is not None:
                    ws.cell(rr, c.column, c.value).number_format = c.number_format
    return ws


# ------------------------------------------------------------------ 主流程
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", required=True)
    ap.add_argument("--template", required=True)
    ap.add_argument("--month", required=True, type=int)
    ap.add_argument("--out", required=True)
    ap.add_argument("--roster", action="append", default=[])
    ap.add_argument("--config")
    ap.add_argument("--target-source", choices=["sheet", "report"], default="sheet")
    ap.add_argument("--subtotal-from", choices=["members", "sheet"], default="members")
    ap.add_argument("--keep-target", action="store_true")
    ap.add_argument("--skip-dept", nargs="*", default=[])
    ap.add_argument("--dept", nargs="*", default=[], help="确认纳入的配置外新部门")
    ap.add_argument("--allow-new-dept", action="store_true")
    ap.add_argument("--skip-person", nargs="*", default=[])
    ap.add_argument("--person-dept", nargs="*", default=[])
    ap.add_argument("--allow-unknown-dept", action="store_true")
    a = ap.parse_args()
    if not (1 <= a.month <= 12):
        raise SystemExit(f"[X] --month 必须是 1-12（收到 {a.month}）。月份越界会导致目标表全部落空、静默算错。")

    cfg = json.loads(json.dumps(DEFAULT_CONFIG))
    if a.config:
        with open(a.config, encoding="utf-8") as f:
            cfg.update(json.load(f))
    cfg["skip_depts"] = list(cfg.get("skip_depts", [])) + list(a.skip_dept)
    cfg["known_depts"] = list(cfg.get("known_depts") or []) + [x.strip() for x in a.dept]
    cfg["skip_people"] = list(cfg.get("skip_people", [])) + [bare(x) for x in a.skip_person]
    for spec in a.person_dept:
        nm, dept = spec.split(":")
        cfg["person_dept"][nm.strip()] = dept.strip()

    persons, dept_seq = load_base(a.base, cfg)
    # 白名单校验：基础表出现「配置外」的部门 -> 必须人工确认，禁止自动纳入核算
    known = set(cfg.get("known_depts") or [])
    if known:
        new_depts = []
        for d in dept_seq:
            d2 = cfg["alias"].get(d, d)
            if d2 not in known and d2 not in new_depts:
                n = sum(1 for p in persons if cfg["alias"].get(p["base_dept"], p["base_dept"]) == d2)
                new_depts.append((d2, n))
        if new_depts and not a.allow_new_dept:
            raise SystemExit(
                "[X] 基础表出现配置外的新部门（禁止静默纳入核算，请确认后处理）：\n    "
                + "\n    ".join(f"{d}（{n} 人）" for d, n in new_depts)
                + "\n    纳入 -> 加 --dept " + " ".join(d for d, _ in new_depts)
                + "\n    跳过 -> 加 --skip-dept " + " ".join(d for d, _ in new_depts)
                + "\n    （确认无误也可加 --allow-new-dept 一次性放行）")
    blocks = derive_blocks(dept_seq, cfg)
    groups, unknown, in_parent, skipped = assign(persons, blocks, cfg)
    # 参与本月核算的人 = 已归入有效部门块的人（跳过部门/父块未拆的都已排除）
    active = [p for v in groups.values() for p in v]

    if unknown and not a.allow_unknown_dept:
        raise SystemExit("[X] 出现未登记部门（禁止猜，请用 --person-dept 或 --skip-dept 指定）：\n    "
                         + "\n    ".join(f"{n} -> {d}" for n, d in unknown))
    if in_parent:
        raise SystemExit("[X] 下列人员仍留在父部门，未指定 1部/2部（禁止猜）：\n    "
                         + "、".join(in_parent) + "\n    请加 --person-dept \"姓名:特定渠道2部\"")

    sheet_tgt, sheet_sums = load_target_sheet(a.template, a.month)
    missing, fallback = resolve_targets(active, sheet_tgt, a.target_source)
    if missing:
        raise SystemExit("[X] 以下人员既不在目标表、报表指标也为 0（禁止推算）：\n    " + "、".join(missing))

    rows = make_rows(blocks, groups, cfg, sheet_sums, a.subtotal_from)

    wb = openpyxl.load_workbook(a.template)
    title = wb.sheetnames[0]
    if not a.keep_target and "目标表" in wb.sheetnames:
        rebuild_target_sheet(wb, a.template, blocks, groups, cfg)
    if a.roster:
        replace_personnel(wb, a.roster)
    ws, last = build_amount_sheet(wb, a.template, rows, title)
    wb.save(a.out)

    P = S = 0
    for row in rows:
        if row[1] == "小计":
            S += row[2]
            print(f"  小计 {row[0]:<10} 目标={row[3]:>11,.0f} 实际={row[4]:>14,.2f} 奖金={row[2]}")
        else:
            P += row[2]
            print(f"  {row[0]:<10} {row[1]:<12} 奖金={row[2]}")
    print("-" * 100)
    print(f"输出: {a.out}")
    print(f"月份={a.month} | 部门块={blocks}")
    print(f"金额表末行={last} | 目标表{'未重建' if a.keep_target else '已重建'} | 参与核算={len(active)} 人 | 人员表={len(a.roster)}份")
    if skipped:
        print(f"[i] 跳过部门（本月不做）{sorted(set(d for _, d in skipped))}：{len(skipped)} 人 -> "
              + "、".join(n for n, _ in skipped))
    if fallback:
        print(f"[i] 目标表查不到、回落到报表指标：{'、'.join(fallback)}")
    print(f"员工个人合计 {P} 元 / 小计合计 {S} 元 / ★ 总奖金 {P + S} 元")
    zero = [p["name"] for p in active if not p["disabled"] and p["T"]["monthly"] == 0]
    if zero:
        print(f"[!] 无月度销售指标的非停用人员: {zero}")


if __name__ == "__main__":
    main()
